# -*- coding: utf-8 -*-
"""
Datenschicht der Arbeitsstände-Anwendung.

Enthält ausschließlich Lade-/Speicherlogik und das Datenmodell -- keine
GUI-Abhängigkeit. Dadurch lässt sich diese Datei vollständig ohne Fenster
testen (siehe test_arbeitsstaende_data.py).
"""
from __future__ import annotations

import json
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from typing import List, Optional

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

STATUS_OPTIONEN = ["Ausstehend", "In Bearbeitung", "Abgeschlossen", "Nicht bestanden", "Sonstiges"]
KURSUNG_OPTIONEN = ["", "E", "G"]

BAUSTEIN_SPALTEN = [
    "name", "status", "bausteinarbeit",
    "lzk_datum_1", "lzk_note_1", "lzk_datum_2", "lzk_note_2",
    "halbjahr", "bemerkung", "lzk_bem_1", "lzk_bem_2",
]

STATUS_FARBEN = {
    "Abgeschlossen": "C6EFCE",
    "In Bearbeitung": "FFEB9C",
    "Nicht bestanden": "FFC7CE",
    "Ausstehend": "F2F2F2",
    "Sonstiges": "D9D2E9",
}


def _to_date(value):
    """Excel/openpyxl liefert Daten manchmal als datetime, manchmal als Text."""
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%m/%d/%y", "%m/%d/%Y"):
            try:
                return datetime.strptime(value.strip(), fmt).date()
            except ValueError:
                continue
    return None


# Halbjahre werden genauso gebildet wie in der Lerntheken-App (server.js,
# halbjahrForDate): Schuljahr laeuft ab August, Format "<StartJJ><EndJJ>_<1|2>".
# Der Januar zaehlt noch zum ERSTEN Halbjahr des im August gestarteten Jahres.
def halbjahr_fuer_datum(d: Optional[date] = None) -> str:
    d = d or date.today()
    m, y = d.month, d.year
    if m >= 8:
        start, sem = y, 1
    elif m == 1:
        start, sem = y - 1, 1
    else:
        start, sem = y - 1, 2
    return f"{str(start)[-2:]}{str(start + 1)[-2:]}_{sem}"


def halbjahr_optionen(zusaetzlich=(), jahre_zurueck: int = 3, jahre_vor: int = 1):
    """Auswahlliste fuer das Halbjahr-Feld, neueste zuerst.

    `zusaetzlich` nimmt Werte auf, die schon in der Datei stehen -- sonst waere
    ein alter oder von Hand gesetzter Wert im Dropdown nicht enthalten und beim
    Bearbeiten still verloren.
    """
    heute = date.today()
    startjahr = heute.year if heute.month >= 8 else heute.year - 1
    optionen = []
    for jahr in range(startjahr - jahre_zurueck, startjahr + jahre_vor + 1):
        for sem in (1, 2):
            optionen.append(f"{str(jahr)[-2:]}{str(jahr + 1)[-2:]}_{sem}")
    for wert in zusaetzlich:
        wert = (wert or "").strip()
        if wert and wert not in optionen:
            optionen.append(wert)
    return sorted(set(optionen), reverse=True)


# Erste Zeile der Bemerkung bei Zeilen, die aus der Lerntheken-App stammen.
# Nur solche Zeilen werden bei einem erneuten Abruf angefasst -- von Hand
# gepflegte Bausteine bleiben dadurch garantiert unberuehrt, selbst wenn sie
# zufaellig genauso heissen.
LT_MARKER = "[Lerntheken-App]"

# Name der Talk-Zeile; das Halbjahr wird angehaengt (z.B. "Mathe-Talks 2627_1").
LT_TALK_NAME = "Mathe-Talks"


def _ist_app_zeile(baustein) -> bool:
    return (baustein.bemerkung or "").lstrip().startswith(LT_MARKER)


def _fmt_kurz(datum) -> str:
    d = _to_date(datum)
    return d.strftime("%d.%m.%Y") if d else ""


def _erledigte_stationen(progress: dict, key: str) -> int:
    """Anzahl erledigter Stationen aus dem Fortschritt einer Lerntheke.

    Der Wert ist ein JSON-Array mit den IDs der erledigten Stationen; Abgabe-
    und Zwischenspeicher-Schluessel gehoeren nicht dazu.
    """
    roh = (progress or {}).get(key)
    if not roh:
        return 0
    if isinstance(roh, str):
        try:
            roh = json.loads(roh)
        except ValueError:
            return 0
    return len(roh) if isinstance(roh, list) else 0


def lt_lerntheke_zeilen(progress: dict, lzk_liste, lerntheken, aktuelles_hj: str):
    """Eine Zeile je Lerntheke, an der tatsaechlich gearbeitet wurde.

    Grundlage ist der GESAMTSTAND (erledigte Stationen aus dem Fortschritt) plus
    die LZK-Eintraege -- nicht das Stations-Ereignislog, das erst seit seiner
    Einfuehrung Daten hat.

    Halbjahr-Zuordnung: das Halbjahr der (frueheren) LZK, sonst das laufende --
    fuer den reinen Stationsfortschritt gibt es kein Datum.
    """
    lzk_je_lerntheke = {}
    for l in lzk_liste or []:
        lzk_je_lerntheke.setdefault(l.get("lerntheke"), []).append(l)

    zeilen = []
    for lt in lerntheken:
        key = lt.get("key")
        if not key:
            continue
        erledigt = _erledigte_stationen(progress, key)
        lzks = lzk_je_lerntheke.get(key, [])
        if not erledigt and not lzks:
            continue                      # nie angefasst -> keine Zeile

        basis = next((l for l in lzks if l.get("typ") == "Basis"), None)
        aufbau = next((l for l in lzks if l.get("typ") == "Aufbau"), None)

        teile = []
        if erledigt:
            gesamt = lt.get("total") or 0
            teile.append(f"{erledigt} von {gesamt} Stationen erledigt" if gesamt
                         else f"{erledigt} Stationen erledigt")
        for l in (basis, aufbau):
            if not l:
                continue
            stand = {"bestanden": "bestanden",
                     "nicht_bestanden": "nicht bestanden"}.get(l.get("status"), "geplant")
            # Wunschtermine von Schueler:innen sind noch keine Termine.
            if l.get("anfrage") == "offen":
                stand = "angefragt"
            elif l.get("anfrage") == "abgelehnt":
                stand = "Anfrage abgelehnt"
            pk = l.get("pokale") or 0
            teile.append(f"{l.get('typ')}-LZK {stand}"
                         + (f", {pk} Flammen" if pk else "")
                         + (f" ({_fmt_kurz(l.get('datum'))})" if l.get("datum") else ""))

        if aufbau and aufbau.get("status") == "bestanden":
            status = "Abgeschlossen"
        elif any(l.get("status") == "nicht_bestanden" for l in lzks):
            status = "Nicht bestanden"
        else:
            status = "In Bearbeitung"

        # Halbjahr aus dem juengsten FESTEN LZK-Datum, sonst das laufende Halbjahr.
        daten = [_to_date(l.get("datum")) for l in lzks if l.get("datum") and not l.get("anfrage")]
        halbjahr = halbjahr_fuer_datum(max(d for d in daten if d)) if any(daten) else aktuelles_hj

        zeilen.append({
            "name": lt.get("title") or key,
            "halbjahr": halbjahr,
            # Nur bestaetigte Termine kommen in die Datumsspalten; ein Wunschtermin
            # laesst ein vorhandenes Datum hier unangetastet.
            "lzk_datum_1": _to_date(basis.get("datum")) if basis and not basis.get("anfrage") else None,
            "lzk_datum_2": _to_date(aufbau.get("datum")) if aufbau and not aufbau.get("anfrage") else None,
            "status": status,
            "bemerkung": " · ".join(teile),
        })
    zeilen.sort(key=lambda z: (z["halbjahr"], z["name"]))
    return zeilen


# Ergebnis einer LZK in EINEM Feld:
#   ""                 hier (noch) nicht bewertet - wird nie zum Server geschickt
#   "nicht_bestanden"
#   "0".."3"           bestanden, mit so vielen Flammen ("0" nur fuer Altbestand)
# Auf dem Server heisst das status + pokale; Flammen bedeuten dort bestanden.
LZK_ERGEBNIS_WERTE = ["", "nicht_bestanden", "0", "1", "2", "3"]
LZK_ERGEBNIS_TEXT = {
    "": "– (nicht bewertet)",
    "nicht_bestanden": "✗ nicht bestanden",
    "0": "✓ bestanden (ohne Flammen)",
    "1": "🔥 bestanden",
    "2": "🔥🔥 bestanden",
    "3": "🔥🔥🔥 bestanden",
}
LZK_ERGEBNIS_SYMBOL = {"nicht_bestanden": "✗", "0": "✓", "1": "🔥", "2": "🔥🔥", "3": "🔥🔥🔥"}


def lzk_ergebnis_von_server(eintrag) -> str:
    """Server-Eintrag (status, pokale) -> Ergebnis-Wert. 'ausstehend' ist ""."""
    if not eintrag:
        return ""
    status = eintrag.get("status") or "ausstehend"
    if status == "nicht_bestanden":
        return "nicht_bestanden"
    if status == "bestanden":
        try:
            pk = int(eintrag.get("pokale") or 0)
        except (TypeError, ValueError):
            pk = 0
        return str(min(3, max(0, pk)))
    return ""


def lzk_ergebnis_zum_server(wert):
    """Ergebnis-Wert -> (status, pokale). None fuer "": nicht bewertet wird nie
    geschickt, sonst wuerde eine online eingetragene Bewertung geloescht."""
    if wert == "nicht_bestanden":
        return ("nicht_bestanden", 0)
    if wert in ("0", "1", "2", "3"):
        return ("bestanden", int(wert))
    return None


# LZK 1 in der Bausteinzeile ist die Basis-, LZK 2 die Aufbau-LZK -- genau so
# baut lt_lerntheke_zeilen() die Zeilen aus den Serverdaten auf.
LZK_TYP_JE_NUMMER = {1: "Basis", 2: "Aufbau"}


def lt_lzk_ergebnis_unterschiede(student, konto: dict, titel_je_key: dict):
    """Ergebnisse (bestanden/Flammen) einer Person: hier gegen den Server.

    Wie bei den Terminen zaehlen nur App-Zeilen, deren Name zu einer Lerntheke
    gehoert. Rueckgabe: je abweichender LZK ein dict mit 'hier' und 'online'
    (Werte aus LZK_ERGEBNIS_WERTE) und allem, was zum Uebertragen noetig ist -
    auch das Datum auf dem Server, damit es beim Schreiben unveraendert bleibt.
    """
    key_je_titel = {(t or "").strip().lower(): k for k, t in (titel_je_key or {}).items()}
    server = {}
    for eintrag in konto.get("lzk") or []:
        server[(eintrag.get("lerntheke"), eintrag.get("typ"))] = eintrag

    unterschiede = []
    for b in student.bausteine:
        if not _ist_app_zeile(b):
            continue
        key = key_je_titel.get((b.name or "").strip().lower())
        if not key:
            continue
        for nummer, typ in LZK_TYP_JE_NUMMER.items():
            hier = getattr(b, f"lzk_ergebnis_{nummer}") or ""
            eintrag = server.get((key, typ))
            if (eintrag or {}).get("anfrage"):
                continue      # eine Anfrage ist noch kein Termin, also auch nicht zu bewerten
            online = lzk_ergebnis_von_server(eintrag)
            if hier == online:
                continue
            unterschiede.append({
                "person": student.voller_name,
                "alias": student.alias.strip().lower(),
                "user_id": konto.get("id"),
                "lerntheke": key,
                "titel": b.name,
                "typ": typ,
                "nummer": nummer,
                "hier": hier,
                "online": online,
                "datum_server": (eintrag or {}).get("datum"),
                "baustein": b,
            })
    return unterschiede


# Freie LZK (ohne Lerntheke) kennt das Tool nur fuer Mathe - wie alles andere hier.
FREIE_LZK_FACH = "mathe"


def _titel_norm(text) -> str:
    """Fuer den Titelvergleich: Gross/klein und Leerzeichen egal."""
    return " ".join(str(text or "").lower().split())


def lt_freie_lzk(konto: dict) -> list:
    """Die freien Mathe-LZK eines Kontos: ohne Lerntheke, fester Termin (keine
    offene oder abgelehnte Anfrage), mit Datum. Eintraege ohne ID stammen von einem
    aelteren Server und lassen sich nicht gezielt aendern - sie bleiben aussen vor."""
    return [e for e in (konto.get("lzk") or [])
            if not e.get("lerntheke") and e.get("id")
            and (e.get("fach") or FREIE_LZK_FACH) == FREIE_LZK_FACH
            and not e.get("anfrage") and e.get("datum")]


def _lzk_online_sauber(wert) -> Optional[dict]:
    """Die gemerkte Verknuepfung eines LZK-Platzes mit einer freien LZK online, geprueft.

    {"id", "datum" ("JJJJ-MM-TT" oder None), "ergebnis"}: der Stand, auf den sich
    Liste und Server beim letzten Abgleich geeinigt haben. Ohne gueltige ID gibt es
    keine Verknuepfung -- dann ordnet wieder der Titel zu.
    """
    if not isinstance(wert, dict):
        return None
    try:
        lzk_id = int(wert.get("id"))
    except (TypeError, ValueError):
        return None
    if lzk_id <= 0:
        return None
    erg = str(wert.get("ergebnis") or "")
    return {"id": lzk_id, "datum": _iso(wert.get("datum")),
            "ergebnis": erg if erg in LZK_ERGEBNIS_WERTE else ""}


def lzk_verknuepfung(eintrag: dict) -> dict:
    """Die Verknuepfung, die sich ein Platz nach einem Abgleich merkt: der Stand online."""
    return {"id": eintrag.get("id"), "datum": _iso(eintrag.get("datum")),
            "ergebnis": lzk_ergebnis_von_server(eintrag)}


def _lzk_seite(link, hier, online):
    """Was seit dem letzten Abgleich geaendert wurde: "hier", "online", "beide" oder
    None. `hier` und `online` sind je (datum, ergebnis)."""
    if not link:
        return None
    basis = (_to_date(link.get("datum")), link.get("ergebnis") or "")
    h, o = hier != basis, online != basis
    return "beide" if h and o else "hier" if h else "online" if o else None


def _lzk_online_titel(eintrag, titel_je_key=None) -> str:
    """Wie eine LZK online heisst: Lerntheke oder Thema."""
    if eintrag.get("lerntheke"):
        return (titel_je_key or {}).get(eintrag["lerntheke"]) or eintrag["lerntheke"]
    return (eintrag.get("thema") or "").strip() or "ohne Titel"


def _freie_lzk_paar(student, e, b, nummer, thema, teil, verknuepft=False) -> dict:
    datum = _to_date(e.get("datum"))
    hier_datum = getattr(b, f"lzk_datum_{nummer}")
    hier_erg = getattr(b, f"lzk_ergebnis_{nummer}") or ""
    online_erg = lzk_ergebnis_von_server(e)
    link = getattr(b, f"lzk_online_{nummer}", None) if verknuepft else None
    return {
        "person": student.voller_name, "alias": student.alias.strip().lower(), "eintrag": e, "id": e.get("id"),
        "baustein": b, "nummer": nummer, "thema": thema, "teil": teil,
        "hier_datum": hier_datum, "online_datum": datum,
        "hier_erg": hier_erg, "online_erg": online_erg,
        "gleich": hier_datum == datum and hier_erg == online_erg,
        "link": link,
        "seite": _lzk_seite(link, (_to_date(hier_datum), hier_erg), (datum, online_erg)),
    }


def lt_freie_lzk_abgleich(student, konto: dict):
    """Ordnet die freien Mathe-LZK einer Person ihren Bausteinen zu.

    Zuerst gilt die Verknuepfung, die sich ein LZK-Platz gemerkt hat (lzk_online_N --
    entsteht, wenn das Tool die LZK anlegt oder abgleicht): die ID entscheidet, auch
    wenn online das Thema oder hier der Name geaendert wurde. Sonst der Titel: Thema
    der LZK == Baustein-Name (Gross/klein und Leerzeichen egal); eine "Aufbau"-LZK
    kommt in LZK-Platz 2, alle anderen in Platz 1. Ziel sind nur VON HAND gepflegte
    Bausteine: die App-Zeilen gehoeren den Lerntheken-LZK, sonst stritten sich zwei
    LZK um denselben Platz.

    Rueckgabe (paare, offen): je zugeordneter LZK ein dict mit Baustein, Platz,
    Datum/Ergebnis hier und online, 'gleich', 'link' (die gemerkte Verknuepfung,
    wenn sie zu genau dieser LZK gehoert) und 'seite' (was seit dem letzten
    Abgleich geaendert wurde); `offen` beschreibt jede LZK, die sich nicht eindeutig
    zuordnen liess. Hier wird nichts veraendert.
    """
    person = student.voller_name
    paare, offen = [], []
    eintraege = lt_freie_lzk(konto)
    online_ids = {e.get("id") for e in eintraege}

    # Verknuepfte Plaetze: welche ID gehoert zu welchem Platz?
    verknuepft, belegt = {}, set()
    for b in student.bausteine:
        if _ist_app_zeile(b):
            continue
        for nummer in (1, 2):
            link = getattr(b, f"lzk_online_{nummer}", None)
            if link and link.get("id") in online_ids:
                verknuepft.setdefault(link["id"], []).append((b, nummer))
                belegt.add((id(b), nummer))

    je_platz = {}
    for e in eintraege:
        datum = _to_date(e.get("datum"))
        typ = e.get("typ") or "LZK"
        wann = _fmt_kurz(datum)
        thema = (e.get("thema") or "").strip()
        teil = f"{typ}-LZK" if typ in ("Basis", "Aufbau") else "LZK"
        plaetze = verknuepft.get(e.get("id"))
        if plaetze:
            if len(plaetze) > 1:
                # Dieselbe ID an zwei Plaetzen (etwa ein kopierter Baustein): nicht raten.
                offen.append(f"{person}: „{thema or 'ohne Titel'}“ ({teil}, {wann}) – "
                             f"mit mehreren Bausteinen verknüpft")
                continue
            b, nummer = plaetze[0]
            paare.append(_freie_lzk_paar(student, e, b, nummer, thema, teil, verknuepft=True))
            continue
        if not thema:
            offen.append(f"{person}: {teil} am {wann} ohne Titel – online ein Thema eintragen")
            continue
        kandidaten = [b for b in student.bausteine
                      if not _ist_app_zeile(b) and _titel_norm(b.name) == _titel_norm(thema)]
        if len(kandidaten) > 1:
            # Gleicher Name in mehreren Halbjahren: das Halbjahr der LZK entscheidet,
            # sonst eine Zeile ohne Halbjahr. Bleibt es mehrdeutig, wird nichts geraten.
            hj = halbjahr_fuer_datum(datum) if datum else ""
            passend = [b for b in kandidaten if b.halbjahr == hj] or [b for b in kandidaten if not b.halbjahr]
            kandidaten = passend if len(passend) == 1 else kandidaten
        if not kandidaten:
            app = any(_ist_app_zeile(b) and _titel_norm(b.name) == _titel_norm(thema)
                      for b in student.bausteine)
            offen.append(f"{person}: „{thema}“ ({teil}, {wann}) – "
                         + ("das ist eine Lerntheken-Zeile; diese LZK bitte in der Lerntheke eintragen"
                            if app else "kein Baustein mit diesem Namen"))
            continue
        if len(kandidaten) > 1:
            offen.append(f"{person}: „{thema}“ ({teil}, {wann}) – mehrere Bausteine mit diesem Namen")
            continue
        b = kandidaten[0]
        nummer = 2 if typ == "Aufbau" else 1
        schluessel = (id(b), nummer)
        if schluessel in belegt:
            offen.append(f"{person}: „{thema}“ ({teil}, {wann}) – LZK {nummer} dieses Bausteins "
                         f"gehört schon zu einer anderen LZK")
            continue
        # Zwei LZK fuer denselben Platz (z.B. nachgeschrieben): die juengste gilt.
        vorher = je_platz.get(schluessel)
        if vorher is not None:
            alt, neu = sorted((vorher, e), key=lambda x: _to_date(x.get("datum")) or date.min)
            offen.append(f"{person}: „{thema}“ ({teil}, {_fmt_kurz(_to_date(alt.get('datum')))}) – "
                         f"für diesen Platz gibt es eine neuere LZK")
            if neu is vorher:
                continue
            paare[:] = [x for x in paare if x["eintrag"] is not vorher]
        je_platz[schluessel] = e
        paare.append(_freie_lzk_paar(student, e, b, nummer, thema, teil))
    return paare, offen


def lt_freie_lzk_senden(student, konto: dict, heute: Optional[date] = None, titel_je_key=None):
    """LZK-Termine aus VON HAND gepflegten Bausteinen, die online fehlen oder hier
    geaendert wurden -- als freie Mathe-LZK: Thema = Baustein-Name, LZK 1 = Basis,
    LZK 2 = Aufbau.

    Nur die Richtung Liste -> Server, und nur, was HIER geaendert wurde: jeder Platz
    merkt sich den zuletzt abgeglichenen Stand (lzk_online_N). Hat sich seither nur
    online etwas getan (verschoben, bewertet), bleibt es online unangetastet und
    wird gemeldet -- hereingeholt wird es beim Abruf (lt_freie_lzk_uebernehmen);
    haben sich beide Seiten verschieden geaendert, entscheidet der Mensch
    ("Freie Mathe-LZK zuordnen…"). Online geloescht wird nie etwas, und ein leeres
    Feld hier ueberschreibt nichts.

    Neu angelegt wird, was ansteht (ab heute) oder hier schon ein Ergebnis hat --
    vergangene Termine ohne Ergebnis nicht, sonst stuenden online auf einen Schlag
    lauter alte LZK "zu bewerten" da. Ebenso wenig, wenn online zu diesem Baustein
    eine Anfrage offen ist oder am selben Tag schon eine Mathe-LZK steht (etwa die
    Lerntheken-LZK): dieselbe LZK stuende sonst zweimal im Kalender.

    Rueckgabe (auftraege, uebersprungen, vergangen):
      auftraege -- dicts mit "art": "anlegen" | "aendern" | "verknuepfen" (nur hier:
        die gemerkte Verknuepfung setzen, nachziehen oder loesen) und "link" (so
        soll sie nach Erfolg lauten)
      uebersprungen -- {"person", "titel", "typ", "grund"} wie bei lt_lzk_aenderungen
      vergangen -- vergangene Termine DIESES Halbjahres ohne Ergebnis, die online
        fehlen und deshalb nicht angelegt werden
    """
    heute = heute or date.today()
    hj_heute = halbjahr_fuer_datum(heute)
    person, alias, user_id = student.voller_name, student.alias.strip().lower(), konto.get("id")
    paare, _offen = lt_freie_lzk_abgleich(student, konto)
    je_platz = {(id(p["baustein"]), p["nummer"]): p for p in paare}
    mathe = [e for e in (konto.get("lzk") or [])
             if (e.get("fach") or FREIE_LZK_FACH) == FREIE_LZK_FACH and e.get("anfrage") != "abgelehnt"]
    # LZK, die schon zu einem Platz dieser Person gehoeren, sind keine Dubletten.
    eigene = {p["id"] for p in paare}
    for b in student.bausteine:
        if not _ist_app_zeile(b):
            eigene |= {l["id"] for l in (b.lzk_online_1, b.lzk_online_2) if l}
    txt = lambda w: LZK_ERGEBNIS_TEXT.get(w, w) if w else "nicht bewertet"
    zuordnen = "über „Freie Mathe-LZK zuordnen…“"
    abruf = "„Ergebnisse abrufen…“ übernimmt das"
    auftraege, uebersprungen, vergangen = [], [], 0

    for b in student.bausteine:
        if _ist_app_zeile(b) or not (b.name or "").strip():
            continue
        for nummer, typ in LZK_TYP_JE_NUMMER.items():
            lok_d = _to_date(getattr(b, f"lzk_datum_{nummer}"))
            lok_e = getattr(b, f"lzk_ergebnis_{nummer}") or ""
            link = getattr(b, f"lzk_online_{nummer}", None)
            kopf = {"person": person, "alias": alias, "user_id": user_id, "baustein": b,
                    "nummer": nummer, "titel": b.name.strip(), "typ": typ}

            def melde(grund):
                uebersprungen.append({"person": person, "titel": b.name, "typ": typ, "grund": grund})

            p = je_platz.get((id(b), nummer))
            if p is not None:
                e, on_d, on_e = p["eintrag"], p["online_datum"], p["online_erg"]
                if p["link"] is None:
                    # Ueber den Titel gefunden, noch nicht verknuepft.
                    if lok_d is None:
                        continue      # hier kein Termin -- uebernehmen geht ueber "zuordnen"
                    if lok_d != on_d:
                        melde(f"online am {_fmt_kurz(on_d)}, hier {_fmt_kurz(lok_d)} – {zuordnen} klären")
                        continue
                    # Derselbe Tag: dieselbe LZK. "Nicht bewertet" ist der gemeinsame Ausgangsstand.
                    link = {"id": e.get("id"), "datum": _iso(lok_d), "ergebnis": ""}
                basis_d, basis_e = _to_date(link.get("datum")), link.get("ergebnis") or ""
                danach, felder = dict(link), {}
                # Datum
                if lok_d != basis_d and on_d == basis_d:
                    if lok_d is None:
                        melde(f"hier kein Termin mehr, online {_fmt_kurz(on_d)} – streichen nur in OSKlar")
                    else:
                        felder["datum"] = lok_d
                        danach["datum"] = _iso(lok_d)
                elif lok_d == basis_d and on_d != basis_d:
                    melde(f"online auf {_fmt_kurz(on_d)} verschoben – {abruf}")
                elif lok_d != basis_d:
                    if lok_d == on_d:
                        danach["datum"] = _iso(lok_d)
                    else:
                        melde(f"hier {_fmt_kurz(lok_d) or 'kein Termin'}, online {_fmt_kurz(on_d)} – "
                              f"beide geändert, {zuordnen} klären")
                # Ergebnis
                if lok_e != basis_e and on_e == basis_e:
                    ziel = lzk_ergebnis_zum_server(lok_e)
                    if ziel is None:
                        melde(f"hier nicht mehr bewertet, online {txt(on_e)} – bleibt")
                    else:
                        felder["status"], felder["pokale"] = ziel
                        danach["ergebnis"] = lok_e
                elif lok_e == basis_e and on_e != basis_e:
                    melde(f"online bewertet ({txt(on_e)}) – {abruf}" if on_e else
                          f"online nicht mehr bewertet – hier bleibt {txt(lok_e)}")
                elif lok_e != basis_e:
                    if lok_e == on_e:
                        danach["ergebnis"] = lok_e
                    else:
                        melde(f"Ergebnis hier {txt(lok_e)}, online {txt(on_e)} – beide geändert, {zuordnen} klären")
                if felder:
                    auftraege.append({**kopf, "art": "aendern", "id": e.get("id"), "felder": felder,
                                      "alt": on_d, "neu": felder.get("datum", on_d),
                                      "alt_erg": on_e, "neu_erg": danach["ergebnis"], "link": danach})
                elif danach != getattr(b, f"lzk_online_{nummer}", None):
                    auftraege.append({**kopf, "art": "verknuepfen", "link": danach})
                continue

            # Online (noch) keine LZK fuer diesen Platz.
            if link:
                if lok_d is None:
                    # Online weg und hier kein Termin: die Verknuepfung hat ausgedient.
                    auftraege.append({**kopf, "art": "verknuepfen", "link": None})
                    continue
                if lok_d == _to_date(link.get("datum")):
                    melde(f"online gelöscht – hier steht noch {_fmt_kurz(lok_d)}")
                    continue
                # Hier neu terminiert: das ist eine neue LZK.
            if lok_d is None:
                if lok_e:
                    melde("Ergebnis ohne LZK-Datum – ohne Datum lässt sich online nichts anlegen")
                continue
            if lok_d < heute and not lok_e:
                if halbjahr_fuer_datum(lok_d) == hj_heute:
                    vergangen += 1
                continue
            anfrage = next((e for e in mathe if not e.get("lerntheke") and e.get("anfrage") == "offen"
                            and _titel_norm(e.get("thema")) == _titel_norm(b.name)
                            and (2 if e.get("typ") == "Aufbau" else 1) == nummer), None)
            if anfrage:
                melde(f"online angefragt ({_fmt_kurz(anfrage.get('datum')) or 'ohne Datum'}) – "
                      f"erst in OSKlar entscheiden")
                continue
            gleicher_tag = next((e for e in mathe if e.get("id") not in eigene
                                 and _to_date(e.get("datum")) == lok_d), None)
            if gleicher_tag:
                wie = ", angefragt" if gleicher_tag.get("anfrage") else ""
                melde(f"am {_fmt_kurz(lok_d)} steht online schon eine Mathe-LZK "
                      f"(„{_lzk_online_titel(gleicher_tag, titel_je_key)}“{wie}) – nicht doppelt angelegt")
                continue
            auftraege.append({**kopf, "art": "anlegen", "alt": None, "neu": lok_d,
                              "alt_erg": "", "neu_erg": lok_e})
    return auftraege, uebersprungen, vergangen


def lt_freie_lzk_uebernehmen(student, konto: dict):
    """Beim Abruf: was an VERKNUEPFTEN freien Mathe-LZK nur online geaendert wurde
    (verschoben, bewertet), kommt in den Baustein -- ohne Rueckfrage.

    Je Feld nur, wenn es HIER seit dem letzten Abgleich unveraendert ist: dann ist
    der Stand online der neuere. Wurde hier ebenfalls geaendert, bleibt beides
    stehen ("Freie Mathe-LZK zuordnen…" entscheidet); nur hier Geaendertes schickt
    das Senden hinaus. Leeres ueberschreibt nichts: nimmt jemand online die
    Bewertung zurueck, bleibt sie hier stehen und wird gemeldet. Nicht verknuepfte
    LZK ordnet der Abruf weiterhin NICHT selbst zu.

    Rueckgabe (uebernommen, gemeldet): Zeilen fuer den Bericht.
    """
    txt = lambda w: LZK_ERGEBNIS_TEXT.get(w, w) if w else "nicht bewertet"
    uebernommen, gemeldet = [], []
    paare, _offen = lt_freie_lzk_abgleich(student, konto)
    for p in paare:
        link = p["link"]
        if not link:
            continue
        b, nummer = p["baustein"], p["nummer"]
        basis_d, basis_e = _to_date(link.get("datum")), link.get("ergebnis") or ""
        lok_d, lok_e = _to_date(p["hier_datum"]), p["hier_erg"]
        on_d, on_e = p["online_datum"], p["online_erg"]
        wo = f"{p['person']} · {b.name} (LZK {nummer})"
        danach, teile = dict(link), []
        if on_d != lok_d:
            if lok_d == basis_d and on_d is not None:
                setattr(b, f"lzk_datum_{nummer}", on_d)
                danach["datum"] = _iso(on_d)
                teile.append(f"Datum {_fmt_kurz(lok_d) or '–'} → {_fmt_kurz(on_d)}")
        elif on_d != basis_d:
            danach["datum"] = _iso(on_d)          # beide Seiten gleich geaendert
        if on_e != lok_e:
            if lok_e == basis_e:
                if on_e:
                    setattr(b, f"lzk_ergebnis_{nummer}", on_e)
                    danach["ergebnis"] = on_e
                    teile.append(f"Ergebnis {txt(lok_e)} → {txt(on_e)}")
                else:
                    gemeldet.append(f"{wo}: online nicht mehr bewertet – hier bleibt {txt(lok_e)}")
        elif on_e != basis_e:
            danach["ergebnis"] = on_e
        if danach != link:
            setattr(b, f"lzk_online_{nummer}", danach)
        if teile:
            uebernommen.append(f"{wo}: " + ", ".join(teile))
    return uebernommen, gemeldet


def lt_lzk_aenderungen(student, konto: dict, titel_je_key: dict):
    """Vergleicht die LZK-Termine einer Person mit dem Stand auf dem Server.

    Betrachtet werden nur Bausteinzeilen, die aus der Lerntheken-App stammen
    (Marker in der Bemerkung) und deren Name zu einer Lerntheke gehoert. Die
    Termine von Hand gepflegter Zeilen gehen als freie Mathe-LZK hinaus, siehe
    lt_freie_lzk_senden().

    Rueckgabe: (aenderungen, uebersprungen). Jede Aenderung enthaelt alles,
    was der Server zum Speichern braucht -- inklusive `status` und `pokale`
    aus dem bestehenden Eintrag: die Schnittstelle schreibt beide Felder immer
    mit, ein Termin allein wuerde eine bestandene LZK also zuruecksetzen.
    """
    key_je_titel = {(t or "").strip().lower(): k for k, t in (titel_je_key or {}).items()}
    server = {}
    for eintrag in konto.get("lzk") or []:
        server[(eintrag.get("lerntheke"), eintrag.get("typ"))] = eintrag

    aenderungen, uebersprungen = [], []
    for b in student.bausteine:
        if not _ist_app_zeile(b):
            continue
        key = key_je_titel.get((b.name or "").strip().lower())
        if not key:
            continue
        for nummer, typ in LZK_TYP_JE_NUMMER.items():
            lokal = _to_date(getattr(b, f"lzk_datum_{nummer}"))
            vorhanden = server.get((key, typ))
            auf_server = _to_date((vorhanden or {}).get("datum"))
            if lokal == auf_server:
                continue
            if (vorhanden or {}).get("anfrage"):
                # Eine Anfrage entscheidet die Lernbegleitung in OSKlar - das Tool
                # ueberschreibt den Wunschtermin nicht (auch nicht automatisch).
                uebersprungen.append({
                    "person": student.voller_name, "titel": b.name, "typ": typ,
                    "grund": f"online angefragt ({_fmt_kurz(auf_server) or 'ohne Datum'}) – erst in OSKlar entscheiden",
                })
                continue
            if lokal is None:
                # Loeschen waere nicht rueckholbar -- das bleibt der
                # Weboberflaeche vorbehalten, hier wird nur berichtet.
                uebersprungen.append({
                    "person": student.voller_name, "titel": b.name, "typ": typ,
                    "grund": f"lokal kein Termin, auf dem Server {_fmt_kurz(auf_server)}",
                })
                continue
            aenderungen.append({
                "person": student.voller_name,
                "alias": student.alias.strip().lower(),
                "user_id": konto.get("id"),
                "lerntheke": key,
                "titel": b.name,
                "typ": typ,
                "alt": auf_server,
                "neu": lokal,
                "status": (vorhanden or {}).get("status") or "ausstehend",
                "pokale": (vorhanden or {}).get("pokale") or 0,
            })
    return aenderungen, uebersprungen


def lt_fb_besuche(by_halbjahr: dict, fach_key: str = "mathe") -> List[date]:
    """Tage, an denen jemand nachweislich im Fachbuero war.

    Gewertet wird nur, was die Lernbegleitung auf "ok" gesetzt hat; kuenftige
    Termine liefert die App ohnehin nicht als erledigt. Diese Tage gehoeren in
    die FB-Besuchsliste der Person (und damit in den Farbindikator der
    Uebersicht) -- nicht in eine Bausteinzeile, wo sie nur als Zahl stuenden.
    """
    tage = set()
    for bucket in (by_halbjahr or {}).values():
        sub = ((bucket or {}).get("bySubject") or {}).get(fach_key) or {}
        for eintrag in sub.get("inputDetails") or []:
            if (eintrag or {}).get("status") != "erledigt":
                continue
            tag = _to_date((eintrag or {}).get("datum"))
            if tag:
                tage.add(tag)
    return sorted(tage)


def lt_talk_zeile(bucket: dict, fach_key: str = "mathe") -> Optional[str]:
    """Bemerkungstext fuer die Talk-Zeile eines Halbjahres, oder None.

    Die Fachbuero-TEILNAHMEN stehen hier bewusst nicht mehr: sie landen als
    echte Termine in der FB-Besuchsliste (siehe lt_fb_besuche). Was hier bleibt,
    hat dort keinen Platz, weil es kein Besuch ist: versaeumte Termine und
    vergangene, die die Lernbegleitung noch nicht eingetragen hat.
    """
    sub = (bucket.get("bySubject") or {}).get(fach_key) or {}
    geh = sub.get("talksPresented", 0) or 0
    zug = sub.get("talksListened", 0) or 0
    fehlt = sub.get("inputMissed", 0) or 0
    offen = sub.get("inputOpen", 0) or 0
    flammen = (sub.get("pokalePresented", 0) or 0) + (sub.get("pokaleListened", 0) or 0)
    if not (geh or zug or fehlt or offen or flammen):
        return None
    teile = []
    if geh:
        teile.append(f"{geh}x gehalten")
    if zug:
        teile.append(f"{zug}x zugehoert")
    if fehlt:
        teile.append(f"{fehlt}x FaBü gefehlt")
    if offen:
        teile.append(f"{offen}x FaBü offen")
    if flammen:
        teile.append(f"{flammen} Flammen")
    return ", ".join(teile)


def lt_zeilen_aktualisieren(student, by_halbjahr: dict, progress: dict,
                            lzk_liste, lerntheken, stand: str = ""):
    """Traegt die App-Ergebnisse ein und meldet (neu, aktualisiert, fb_besuche).

    Je bearbeiteter Lerntheke eine Bausteinzeile (Gesamtstand + LZK) und je
    Halbjahr eine fuer die Talks. Fachbuero-TEILNAHMEN werden dagegen als
    Besuchstage eingetragen -- sie gehoeren in die FB-Besuchsliste, nicht in
    eine Bausteinzeile.

    Angefasst werden ausschliesslich Zeilen mit LT_MARKER in der Bemerkung.
    Von Hand gepflegte Bausteine bleiben unberuehrt -- auch namensgleiche.
    Noten und ein manuell geaenderter Status werden nie ueberschrieben: Noten
    traegt die Lernbegleitung ein, die App kennt keine.
    """
    neu = aktualisiert = 0

    def _setze(name, halbjahr, bemerkung, status=None, lzk1=None, lzk2=None):
        nonlocal neu, aktualisiert
        text = f"{LT_MARKER} {bemerkung}" + (f" (Stand {stand})" if stand else "")
        vorhanden = next((b for b in student.bausteine
                          if b.name == name and b.halbjahr == halbjahr
                          and _ist_app_zeile(b)), None)
        if vorhanden:
            vorhanden.bemerkung = text
            if lzk1 is not None:
                vorhanden.lzk_datum_1 = lzk1
            if lzk2 is not None:
                vorhanden.lzk_datum_2 = lzk2
            aktualisiert += 1
        else:
            student.bausteine.append(Baustein(
                name=name, status=status or "Sonstiges", halbjahr=halbjahr,
                lzk_datum_1=lzk1, lzk_datum_2=lzk2, bemerkung=text))
            neu += 1

    aktuelles_hj = halbjahr_fuer_datum()
    for z in lt_lerntheke_zeilen(progress, lzk_liste, lerntheken, aktuelles_hj):
        _setze(z["name"], z["halbjahr"], z["bemerkung"], z["status"],
               z["lzk_datum_1"], z["lzk_datum_2"])

    for hj in sorted((by_halbjahr or {}).keys()):
        talk = lt_talk_zeile(by_halbjahr[hj] or {})
        if talk:
            _setze(f"{LT_TALK_NAME} {hj}", hj, talk, "Sonstiges")

    # Fachbuero-Teilnahmen aus der App in die Besuchsliste uebernehmen. Von Hand
    # eingetragene Besuche bleiben erhalten: fb_besuch_eintragen() fuegt nur
    # hinzu, was noch nicht dasteht.
    besuche_neu = 0
    for tag in lt_fb_besuche(by_halbjahr):
        if tag not in student.fb_besuche:
            student.fb_besuch_eintragen(tag)
            besuche_neu += 1

    return neu, aktualisiert, besuche_neu


def alias_vorschlag(vorname: str, nachname: str) -> str:
    """Vorschlag fuer den Lerntheken-Alias: erste 2 Buchstaben des Vornamens,
    Punkt, erste 2 Buchstaben des Nachnamens -- "Anton Berger" -> "an.be".

    Nur ein VORSCHLAG: bei Namensgleichheit muss von Hand nachgebessert werden.
    """
    ersetzungen = {"ä": "ae", "ö": "oe", "ü": "ue", "ß": "ss",
                   "Ä": "Ae", "Ö": "Oe", "Ü": "Ue"}
    def norm(t):
        for a, b in ersetzungen.items():
            t = t.replace(a, b)
        t = unicodedata.normalize("NFKD", t)
        t = "".join(z for z in t if not unicodedata.combining(z))
        return "".join(c for c in t.strip().lower() if c.isalnum())
    v, n = norm(vorname), norm(nachname)
    if not v or not n:
        return ""
    return f"{v[:2]}.{n[:2]}"


def _clean(value):
    if value is None:
        return ""
    return str(value).strip()


# Eigenes Speicherformat. Excel bleibt als Import- und Exportweg erhalten
# (Weitergeben, Ausdrucken), gearbeitet wird aber mit JSON: verlustfrei,
# von Hand lesbar und nachbearbeitbar, und neue Felder brechen alte Dateien
# nicht -- unbekannte Schluessel werden beim Laden schlicht ignoriert.
JSON_FORMAT = "osk-arbeitsstaende"
JSON_VERSION = 1
JSON_ENDUNGEN = (".json",)


def ist_json_pfad(pfad: str) -> bool:
    return str(pfad).lower().endswith(JSON_ENDUNGEN)


def _iso(d) -> Optional[str]:
    d = _to_date(d)
    return d.isoformat() if d else None


def _ohne_leere(d: dict) -> dict:
    """Leere Felder gar nicht erst schreiben -- eine von Hand gelesene Datei
    soll nur zeigen, was tatsaechlich gepflegt ist."""
    return {k: v for k, v in d.items() if v not in (None, "", [], {})}


# ----------------------------------------------------------------------
# Mathe-Talks (2026-09-25)
# ----------------------------------------------------------------------
# Jede Person fuehrt ihre Talks als eigene Liste: einen Eintrag je Talk, den
# sie gehalten (auch zu zweit), mit vorgetragen oder zugehoert hat. Abgerufen
# wird alles, was online steht; zurueckgeschickt werden nur BEWERTUNGEN (und
# bei gehaltenen Talks das Thema) -- und nur, was hier seit dem letzten
# Abgleich geaendert wurde. Grundlage dafuer ist `online`: der Stand, auf den
# sich Liste und Server zuletzt geeinigt haben (wie lzk_online bei den LZK).

TALK_ROLLEN = ("gehalten", "mitvortrag", "zugehoert")
TALK_ROLLEN_TEXT = {"gehalten": "gehalten", "mitvortrag": "mit vorgetragen",
                    "zugehoert": "zugehört"}
# Dieselben Werte wie auf dem Server (TALKING_STATUS_VALUES). "ausstehend" heisst
# "noch nicht bewertet" und wird nie hochgeschickt.
TALK_STATUS_WERTE = ["ausstehend", "erledigt", "nicht_erledigt"]
TALK_STATUS_TEXT = {"ausstehend": "noch nicht bewertet",
                    "erledigt": "✓ ok", "nicht_erledigt": "✗ nicht ok"}
# Auswahl wie auf dem Server (TALKING_QUALITY_EMOJIS); "" = kein Emoji.
TALK_EMOJIS = ["", "🤩", "🌟", "👍", "🙂", "🤔", "💡", "🎯", "🔥"]
TALK_BEWERTUNG = ("status", "flammen", "emoji")


def _talk_datum(wert) -> Optional[date]:
    """Datum eines Talks aus der Server-Antwort.

    Der Server schickt DATE-Spalten als Zeitstempel der lokalen Mitternacht in
    UTC ("2026-09-24T22:00:00.000Z" fuer den 25.09.). Das blosse Abschneiden
    ergaebe den Vortag; +12 Stunden landen fuer jede Zeitzone zwischen -12 und
    +12 Stunden sicher auf dem richtigen Tag.
    """
    if isinstance(wert, str) and "T" in wert:
        try:
            zeit = datetime.fromisoformat(wert.replace("Z", "+00:00"))
        except ValueError:
            return _to_date(wert[:10])
        return (zeit + timedelta(hours=12)).date()
    return _to_date(wert)


def _flammen(wert, hoechstens: int) -> int:
    try:
        n = int(wert or 0)
    except (TypeError, ValueError):
        n = 0
    return max(0, min(hoechstens, n))


@dataclass
class MatheTalk:
    rolle: str = "zugehoert"           # siehe TALK_ROLLEN
    # Kennung online: bei "gehalten" die Buchung (talking_sessions.id), sonst die
    # Einladung (talking_invitations.id) -- daran haengt jeweils die Bewertung.
    online_id: Optional[int] = None
    session_id: Optional[int] = None
    datum: Optional[date] = None
    uhrzeit: str = ""
    halbjahr: str = ""
    thema: str = ""
    # Wer sonst vortraegt ("gehalten"/"mitvortrag") bzw. bei wem man
    # zugehoert hat ("zugehoert").
    mit: str = ""
    status: str = "ausstehend"
    flammen: int = 0
    emoji: str = ""
    bemerkung: str = ""                # nur hier, geht nie zum Server
    online: Optional[dict] = None      # {"status","flammen","emoji","thema"}
    # Beim letzten Abruf online nicht mehr gefunden (Termin geloescht, Person
    # ausgeladen). Wird gemeldet, nie von selbst geloescht.
    online_fehlt: bool = False

    @property
    def max_flammen(self) -> int:
        # Zuhoeren bis 2 Flammen, Vortrag und Mit-Vortrag bis 3 (wie der Server).
        return 2 if self.rolle == "zugehoert" else 3

    @property
    def schluessel(self):
        return ("s" if self.rolle == "gehalten" else "i", self.online_id)

    def bewertung(self) -> dict:
        return {"status": self.status, "flammen": int(self.flammen or 0),
                "emoji": self.emoji or ""}


def talk_als_dict(t: MatheTalk) -> dict:
    return _ohne_leere({
        "rolle": t.rolle, "online_id": t.online_id, "session_id": t.session_id,
        "datum": _iso(t.datum), "uhrzeit": t.uhrzeit, "halbjahr": t.halbjahr,
        "thema": t.thema, "mit": t.mit, "status": t.status,
        "flammen": t.flammen or None, "emoji": t.emoji, "bemerkung": t.bemerkung,
        "online": t.online, "online_fehlt": t.online_fehlt or None,
    })


def talk_aus_dict(d: dict) -> MatheTalk:
    rolle = d.get("rolle") if d.get("rolle") in TALK_ROLLEN else "zugehoert"
    status = d.get("status") if d.get("status") in TALK_STATUS_WERTE else "ausstehend"
    t = MatheTalk(rolle=rolle, online_id=d.get("online_id"), session_id=d.get("session_id"),
                  datum=_to_date(d.get("datum")), uhrzeit=d.get("uhrzeit") or "",
                  halbjahr=d.get("halbjahr") or "", thema=d.get("thema") or "",
                  mit=d.get("mit") or "", status=status,
                  emoji=d.get("emoji") if d.get("emoji") in TALK_EMOJIS else "",
                  bemerkung=d.get("bemerkung") or "",
                  online=d.get("online") if isinstance(d.get("online"), dict) else None,
                  online_fehlt=bool(d.get("online_fehlt")))
    t.flammen = _flammen(d.get("flammen"), t.max_flammen)
    return t


def _dabei(eintrag: dict) -> bool:
    """Zaehlt eine Einladung? Zugesagt, oder schon bewertet (wie der Server)."""
    return (eintrag.get("status") == "angenommen"
            or (eintrag.get("attendedStatus") or "ausstehend") != "ausstehend")


def lt_mathe_talks(slots: list) -> dict:
    """Server-Antwort (/api/admin/talking-slots, Mathe-Talks) -> je Benutzername
    die Liste seiner Talks (dicts in MatheTalk-Form, dazu "online")."""
    je_person = {}

    def dazu(name, **werte):
        name = (name or "").strip().lower()
        if not name:
            return
        werte["online"] = {"status": werte["status"], "flammen": werte["flammen"],
                           "emoji": werte["emoji"], "thema": werte["thema"]}
        je_person.setdefault(name, []).append(werte)

    for slot in slots or []:
        if not slot.get("session_id"):
            continue                       # freier Termin, noch niemand gebucht
        if (slot.get("typ") or "talk") != "talk":
            continue
        datum = _talk_datum(slot.get("datum"))
        gemeinsam = dict(session_id=slot["session_id"], datum=datum,
                         uhrzeit=slot.get("uhrzeit") or "",
                         halbjahr=(slot.get("halbjahr") or "").strip()
                         or (halbjahr_fuer_datum(datum) if datum else ""),
                         thema=slot.get("thema") or "")
        vortrag = slot.get("presenter_username") or ""
        mit_vortrag = [c for c in slot.get("coPresenters") or [] if _dabei(c)]
        mit_namen = [c.get("username") or "" for c in mit_vortrag]

        dazu(vortrag, rolle="gehalten", online_id=slot["session_id"],
             mit=" & ".join(mit_namen),
             status=slot.get("presentedStatus") or "ausstehend",
             flammen=_flammen(slot.get("pokale"), 3),
             emoji=slot.get("qualityEmoji") or "", **gemeinsam)
        for c in mit_vortrag:
            andere = [vortrag] + [n for n in mit_namen if n != c.get("username")]
            dazu(c.get("username"), rolle="mitvortrag", online_id=c.get("id"),
                 mit=" & ".join(n for n in andere if n),
                 status=c.get("attendedStatus") or "ausstehend",
                 flammen=_flammen(c.get("pokale"), 3),
                 emoji=c.get("qualityEmoji") or "", **gemeinsam)
        alle_vortragenden = " & ".join(n for n in [vortrag] + mit_namen if n)
        for i in slot.get("invitees") or []:
            if not _dabei(i):
                continue
            dazu(i.get("username"), rolle="zugehoert", online_id=i.get("id"),
                 mit=alle_vortragenden,
                 status=i.get("attendedStatus") or "ausstehend",
                 flammen=_flammen(i.get("pokale"), 2),
                 emoji=i.get("qualityEmoji") or "", **gemeinsam)
    return je_person


def _talk_text(student, t) -> str:
    return (f"{student.voller_name}: {_fmt_kurz(t.datum)} „{t.thema or 'Talk'}“ "
            f"({TALK_ROLLEN_TEXT.get(t.rolle, t.rolle)})")


def lt_talks_uebernehmen(student, eintraege: list) -> dict:
    """Traegt den Serverstand in student.talks ein -- Drei-Wege-Abgleich.

    Termin-Angaben (Datum, Uhrzeit, wer dabei ist) kommen immer vom Server. Bei
    Bewertung und Thema entscheidet der letzte gemeinsame Stand (`online`):
    hier unveraendert -> Server uebernehmen; nur hier geaendert -> bleibt und
    wartet aufs Hochladen; beide verschieden geaendert -> bleibt hier, wird
    gemeldet. Was online fehlt, wird markiert, nie geloescht.
    """
    bericht = {"neu": 0, "uebernommen": [], "konflikte": [], "fehlen": []}
    vorhanden = {t.schluessel: t for t in student.talks if t.online_id is not None}
    gesehen = set()

    for e in eintraege or []:
        schluessel = ("s" if e["rolle"] == "gehalten" else "i", e["online_id"])
        gesehen.add(schluessel)
        srv = e["online"]
        t = vorhanden.get(schluessel)
        if t is None:
            t = MatheTalk(
                rolle=e["rolle"], online_id=e["online_id"], session_id=e["session_id"],
                datum=e["datum"], uhrzeit=e["uhrzeit"], halbjahr=e["halbjahr"],
                thema=e["thema"], mit=e["mit"], status=srv["status"],
                flammen=srv["flammen"], emoji=srv["emoji"], online=dict(srv))
            student.talks.append(t)
            vorhanden[schluessel] = t
            bericht["neu"] += 1
            continue
        t.session_id, t.datum, t.uhrzeit = e["session_id"], e["datum"], e["uhrzeit"]
        t.halbjahr, t.mit, t.online_fehlt = e["halbjahr"], e["mit"], False
        alt = t.online or {}
        srv_bew = {k: srv[k] for k in TALK_BEWERTUNG}
        alt_bew = {k: alt.get(k) for k in TALK_BEWERTUNG}
        hier_bew = t.bewertung()
        if hier_bew == alt_bew or hier_bew == srv_bew:
            if hier_bew != srv_bew:
                bericht["uebernommen"].append(_talk_text(student, t) + " -- Bewertung")
            t.status, t.flammen, t.emoji = srv["status"], srv["flammen"], srv["emoji"]
            bew_neu = srv_bew
        elif srv_bew == alt_bew:
            bew_neu = alt_bew              # nur hier geaendert: wartet aufs Hochladen
        else:
            bericht["konflikte"].append(_talk_text(student, t)
                                        + " -- hier und online verschieden bewertet, hier bleibt")
            bew_neu = srv_bew              # beim Hochladen sichtbar als "online inzwischen"
        # Das Thema pflegt hier nur der gehaltene Talk; alle anderen folgen dem Server.
        if t.rolle != "gehalten" or t.thema in (alt.get("thema"), srv["thema"]):
            if t.rolle == "gehalten" and t.thema != srv["thema"]:
                bericht["uebernommen"].append(_talk_text(student, t) + " -- Thema")
            t.thema = srv["thema"]
            thema_neu = srv["thema"]
        elif srv["thema"] == alt.get("thema"):
            thema_neu = alt.get("thema")
        else:
            bericht["konflikte"].append(_talk_text(student, t)
                                        + " -- Thema hier und online geaendert, hier bleibt")
            thema_neu = srv["thema"]
        t.online = dict(bew_neu, thema=thema_neu)

    for t in student.talks:
        if t.online_id is not None and t.schluessel not in gesehen and not t.online_fehlt:
            t.online_fehlt = True
            bericht["fehlen"].append(_talk_text(student, t))
    student.talks.sort(key=lambda t: (t.datum or date.min, t.uhrzeit), reverse=True)
    return bericht


def lt_talks_hochladen_liste(student) -> list:
    """Was von dieser Person zum Server geht: [(talk, was, hier, online)].

    `was` ist "bewertung" oder "thema". Nur hier Geaendertes; "noch nicht
    bewertet" und ein leeres Thema gehen nie hoch (leer loescht online nichts).
    """
    return _talk_auftraege(student.talks)


def _talk_auftraege(talks) -> list:
    auftraege = []
    for t in talks:
        if t.online_id is None or t.online_fehlt:
            continue
        alt = t.online or {}
        alt_bew = {k: alt.get(k) for k in TALK_BEWERTUNG}
        if t.status != "ausstehend" and t.bewertung() != alt_bew:
            auftraege.append((t, "bewertung", t.bewertung(), alt_bew))
        if (t.rolle == "gehalten" and (t.thema or "").strip()
                and t.thema.strip() != (alt.get("thema") or "").strip()):
            auftraege.append((t, "thema", t.thema.strip(), alt.get("thema") or ""))
    return auftraege


def talk_wartet(t: MatheTalk) -> bool:
    """Hier geaendert und noch nicht hochgeladen?"""
    return bool(_talk_auftraege([t]))


def talk_hochgeladen(t: MatheTalk, was: str):
    """Nach erfolgreichem Hochladen: der hiesige Stand ist jetzt der gemeinsame."""
    stand = dict(t.online or {})
    if was == "bewertung":
        stand.update(t.bewertung())
    else:
        stand["thema"] = t.thema.strip()
    t.online = stand


def talks_je_halbjahr(student) -> dict:
    """{halbjahr: [talks]} -- neueste Halbjahre zuerst."""
    gruppen = {}
    for t in student.talks:
        gruppen.setdefault(t.halbjahr or "ohne Halbjahr", []).append(t)
    return dict(sorted(gruppen.items(), key=lambda kv: kv[0], reverse=True))


def talks_zusammenfassung(talks: list) -> str:
    geh = sum(1 for t in talks if t.rolle in ("gehalten", "mitvortrag"))
    zu = sum(1 for t in talks if t.rolle == "zugehoert")
    flammen = sum(t.flammen or 0 for t in talks if t.status == "erledigt")
    offen = sum(1 for t in talks if t.status == "ausstehend")
    teile = [f"{geh}x gehalten", f"{zu}x zugehört"]
    if flammen:
        teile.append(f"{flammen} 🔥")
    if offen:
        teile.append(f"{offen} unbewertet")
    return " · ".join(teile)


@dataclass
class Baustein:
    name: str = ""
    status: str = "Ausstehend"
    bausteinarbeit: str = ""
    lzk_datum_1: Optional[date] = None
    lzk_note_1: str = ""
    lzk_datum_2: Optional[date] = None
    lzk_note_2: str = ""
    halbjahr: str = ""
    bemerkung: str = ""
    # Notiz zur jeweiligen LZK (z.B. "nur Teil 1 geschrieben", "Nachschreibtermin")
    lzk_bem_1: str = ""
    lzk_bem_2: str = ""
    # Ergebnis der jeweiligen LZK, wie es auch auf dem Server steht (siehe
    # LZK_ERGEBNIS_WERTE). Die "LZK-Note" daneben ist Freitext und bleibt lokal.
    lzk_ergebnis_1: str = ""
    lzk_ergebnis_2: str = ""
    # Verknuepfung mit der freien Mathe-LZK online (nur von Hand gepflegte
    # Bausteine): {"id", "datum", "ergebnis"} -- der Stand, auf den sich Liste und
    # Server beim letzten Abgleich geeinigt haben. Daran erkennt das Senden, WELCHE
    # Seite seither geaendert wurde (siehe lt_freie_lzk_senden).
    lzk_online_1: Optional[dict] = None
    lzk_online_2: Optional[dict] = None


@dataclass
class Student:
    vorname: str
    nachname: str
    # Benutzername in der Lerntheken-App ("Lerntheken-Alias"). Wird im
    # Namen-Blatt gepflegt und ist die verbindliche Zuordnung zwischen dieser
    # Liste und der App -- bewusst NICHT jedes Mal aus dem Namen abgeleitet,
    # weil zwei Personen denselben Kurznamen ergeben koennen (Anton Berger /
    # Anna Bergmann -> beide "an.be") und Namen sich aendern.
    alias: str = ""
    kursung: str = ""
    jahrgangsstufe: Optional[int] = None
    letzter_besuch_fb: Optional[date] = None
    hj_note: str = ""
    # Frei setzbare Frist neben den LZK-Terminen (Referat abgeben, Material
    # drucken, ...). Die Bemerkung ist zugleich der Anlass, der in der Uebersicht
    # neben dem Datum steht.
    sonstige_deadline: Optional[date] = None
    sonstige_deadline_bemerkung: str = ""
    bausteine: List[Baustein] = field(default_factory=list)
    # Name des Original-Arbeitsblatts, falls aus einer Datei geladen.
    # Wird für gezieltes Aktualisieren beim Speichern gebraucht.
    _quellblatt: Optional[str] = None
    # Volle FB-Besuchshistorie (ein Eintrag pro Tag). Wird aus dem
    # Anwesenheits-Raster im Namen-Blatt gelesen/geschrieben und ist die
    # Grundlage für "letzter_besuch_fb" sowie für den Schnell-Eintrag
    # ("War heute im FB" / "Nachtragen") in der Oberfläche.
    fb_besuche: List[date] = field(default_factory=list)
    # Mathe-Talks der Person (gehalten, mit vorgetragen, zugehoert) -- nur im
    # JSON-Format, Excel kennt sie nicht. Siehe MatheTalk.
    talks: List[MatheTalk] = field(default_factory=list)

    def fb_besuch_eintragen(self, tag: date):
        if tag not in self.fb_besuche:
            self.fb_besuche.append(tag)
            self.fb_besuche.sort()
        if self.letzter_besuch_fb is None or tag > self.letzter_besuch_fb:
            self.letzter_besuch_fb = tag

    @property
    def voller_name(self) -> str:
        return f"{self.vorname} {self.nachname}".strip()


def baustein_als_dict(b: Baustein) -> dict:
    daten = {
        "name": b.name,
        "status": b.status,
        "bausteinarbeit": b.bausteinarbeit,
        "halbjahr": b.halbjahr,
        "bemerkung": b.bemerkung,
    }
    for nr, (datum, note, bem, erg, online) in enumerate(
            ((b.lzk_datum_1, b.lzk_note_1, b.lzk_bem_1, b.lzk_ergebnis_1, b.lzk_online_1),
             (b.lzk_datum_2, b.lzk_note_2, b.lzk_bem_2, b.lzk_ergebnis_2, b.lzk_online_2)), start=1):
        lzk = _ohne_leere({"datum": _iso(datum), "note": note, "bemerkung": bem, "ergebnis": erg,
                           "online": _lzk_online_sauber(online)})
        if lzk:
            daten[f"lzk_{nr}"] = lzk
    return _ohne_leere(daten)


def baustein_aus_dict(d: dict) -> Baustein:
    b = Baustein(
        name=d.get("name") or "",
        status=d.get("status") or "Ausstehend",
        bausteinarbeit=d.get("bausteinarbeit") or "",
        halbjahr=d.get("halbjahr") or "",
        bemerkung=d.get("bemerkung") or "",
    )
    for nr, felder in ((1, ("lzk_datum_1", "lzk_note_1", "lzk_bem_1", "lzk_ergebnis_1")),
                       (2, ("lzk_datum_2", "lzk_note_2", "lzk_bem_2", "lzk_ergebnis_2"))):
        lzk = d.get(f"lzk_{nr}") or {}
        setattr(b, felder[0], _to_date(lzk.get("datum")))
        setattr(b, felder[1], lzk.get("note") or "")
        setattr(b, felder[2], lzk.get("bemerkung") or "")
        # Aeltere Dateien kennen kein Ergebnis; ein unbekannter Wert wird nicht
        # geraten, sondern bleibt leer ("nicht bewertet").
        erg = str(lzk.get("ergebnis") or "")
        setattr(b, felder[3], erg if erg in LZK_ERGEBNIS_WERTE else "")
        setattr(b, f"lzk_online_{nr}", _lzk_online_sauber(lzk.get("online")))
    return b


def student_als_dict(s: "Student") -> dict:
    return _ohne_leere({
        "vorname": s.vorname,
        "nachname": s.nachname,
        "lerntheken_alias": s.alias,
        "kursung": s.kursung,
        "jahrgangsstufe": s.jahrgangsstufe,
        "hj_note": s.hj_note,
        "sonstige_deadline": _iso(s.sonstige_deadline),
        "sonstige_deadline_anlass": s.sonstige_deadline_bemerkung,
        "fb_besuche": [d.isoformat() for d in s.fb_besuche],
        "bausteine": [baustein_als_dict(b) for b in s.bausteine],
        "mathe_talks": [talk_als_dict(t) for t in s.talks],
    })


def student_aus_dict(d: dict) -> "Student":
    jahrgang = d.get("jahrgangsstufe")
    try:
        jahrgang = int(jahrgang) if jahrgang not in (None, "") else None
    except (TypeError, ValueError):
        jahrgang = None
    s = Student(
        vorname=d.get("vorname") or "",
        nachname=d.get("nachname") or "",
        alias=(d.get("lerntheken_alias") or "").strip().lower(),
        kursung=d.get("kursung") or "",
        jahrgangsstufe=jahrgang,
        hj_note=d.get("hj_note") or "",
        sonstige_deadline=_to_date(d.get("sonstige_deadline")),
        sonstige_deadline_bemerkung=d.get("sonstige_deadline_anlass") or "",
        bausteine=[baustein_aus_dict(b) for b in d.get("bausteine") or []],
    )
    besuche = sorted({b for b in (_to_date(x) for x in d.get("fb_besuche") or []) if b})
    s.fb_besuche = besuche
    s.letzter_besuch_fb = besuche[-1] if besuche else None
    s.talks = [talk_aus_dict(t) for t in d.get("mathe_talks") or [] if isinstance(t, dict)]
    return s


class Arbeitsstaende:
    """Hält den kompletten Datenbestand und kapselt Laden/Speichern."""

    def __init__(self):
        self.students: List[Student] = []
        self.vorlage_bausteine: List[str] = []
        self.dateipfad: Optional[str] = None
        self._wb: Optional[Workbook] = None  # zuletzt geladene/erzeugte Mappe

    # ------------------------------------------------------------------
    # Eigenes Format (JSON)
    # ------------------------------------------------------------------
    def als_dict(self) -> dict:
        return {
            "format": JSON_FORMAT,
            "version": JSON_VERSION,
            "gespeichert_am": date.today().isoformat(),
            "vorlage_bausteine": list(self.vorlage_bausteine),
            "personen": [student_als_dict(s) for s in self.students],
        }

    def aus_dict(self, daten: dict) -> List[str]:
        warnungen: List[str] = []
        if not isinstance(daten, dict):
            raise ValueError("Datei enthaelt kein Arbeitsstaende-Objekt.")
        if daten.get("format") not in (None, JSON_FORMAT):
            warnungen.append(f"Unbekanntes Format '{daten.get('format')}' -- es wird "
                             f"trotzdem versucht, die Daten zu lesen.")
        version = daten.get("version")
        if isinstance(version, int) and version > JSON_VERSION:
            warnungen.append(f"Die Datei stammt aus einer neueren Version (v{version}); "
                             f"unbekannte Angaben gehen beim Speichern verloren.")
        self.vorlage_bausteine = [str(n) for n in daten.get("vorlage_bausteine") or []]
        self.students = []
        for eintrag in daten.get("personen") or []:
            person = student_aus_dict(eintrag)
            if not person.voller_name:
                warnungen.append("Ein Eintrag ohne Namen wurde uebersprungen.")
                continue
            self.students.append(person)
        return warnungen

    def laden_json(self, pfad: str) -> List[str]:
        with open(pfad, encoding="utf-8") as f:
            daten = json.load(f)
        warnungen = self.aus_dict(daten)
        # Eine frisch geladene JSON-Datei hat keine Excel-Mappe im Ruecken; ein
        # spaeterer Excel-Export baut daher eine neue Mappe auf.
        self._wb = None
        self.dateipfad = pfad
        return warnungen

    def speichern_json(self, pfad: str):
        with open(pfad, "w", encoding="utf-8") as f:
            json.dump(self.als_dict(), f, ensure_ascii=False, indent=2)
            f.write("\n")   # abschliessender Zeilenumbruch, git-freundlich
        self.dateipfad = pfad

    # ------------------------------------------------------------------
    # Laden / Speichern -- Format ergibt sich aus der Dateiendung
    # ------------------------------------------------------------------
    def laden(self, pfad: str) -> List[str]:
        """Laedt eine Arbeitsstaende-Datei: .json im eigenen Format, alles
        andere als Excel-Mappe. Gibt eine Liste von Warnungen zurueck."""
        if ist_json_pfad(pfad):
            return self.laden_json(pfad)
        return self.laden_excel(pfad)

    def speichern(self, pfad: str):
        if ist_json_pfad(pfad):
            self.speichern_json(pfad)
        else:
            self.speichern_excel(pfad)

    # ------------------------------------------------------------------
    # Excel (Import und Export)
    # ------------------------------------------------------------------
    def laden_excel(self, pfad: str) -> List[str]:
        """Liest eine Arbeitsstände-Mappe (.xlsx/.xlsm) ein. Gibt eine Liste
        von Warnungen zurück (z.B. übersprungene Blätter)."""
        warnungen: List[str] = []
        wb = load_workbook(pfad, data_only=False)
        self._wb = wb
        self.dateipfad = pfad

        vorlage_name = next(
            (n for n in wb.sheetnames if n.lower().startswith("vorlage")), None
        )
        namen_sheet = "Namen" if "Namen" in wb.sheetnames else None

        schueler_bloetter = [
            n for n in wb.sheetnames
            if n != namen_sheet and n != vorlage_name
        ]

        self.vorlage_bausteine = []
        if vorlage_name:
            self.vorlage_bausteine = self._lies_baustein_namen(wb[vorlage_name])
        else:
            warnungen.append("Kein Vorlage-Blatt gefunden -- Standard-Bausteinliste ist leer.")

        self.students = []
        besuche_by_row = {}
        alias_by_name = {}
        if namen_sheet:
            ws_namen = wb[namen_sheet]
            letzte_spalte = ws_namen.max_column
            # Altes Layout: Raster ab Spalte H. Neues Layout: H = "Anlass",
            # Raster ab Spalte I. Unterscheidbar am Inhalt von H1.
            raster_start = 8 if _to_date(ws_namen.cell(row=1, column=8).value) else 9
            if letzte_spalte >= raster_start:
                spalten_daten = [_to_date(ws_namen.cell(row=1, column=c).value)
                                  for c in range(raster_start, letzte_spalte + 1)]
                for row in range(2, ws_namen.max_row + 1):
                    vn = _clean(ws_namen.cell(row=row, column=1).value)
                    nn = _clean(ws_namen.cell(row=row, column=2).value)
                    if not vn and not nn:
                        continue
                    besuche = []
                    for j, spaltendatum in enumerate(spalten_daten):
                        if spaltendatum is None:
                            continue
                        wert = ws_namen.cell(row=row, column=raster_start + j).value
                        if wert in (1, "1", True):
                            besuche.append(spaltendatum)
                    besuche_by_row[row] = sorted(besuche)

            # Lerntheken-Alias (Spalte C) namensgebunden einlesen -- bewusst in
            # einer EIGENEN Schleife: das Anwesenheitsraster oben laeuft nur,
            # wenn es ueberhaupt Besuchsspalten gibt (max_column >= 8). Haengte
            # der Alias dort mit drin, ginge er in einer Datei ohne Besuche
            # still verloren. Spalte C war bisher ungenutzt, das Raster ab
            # Spalte H verschiebt sich also nicht.
            for row in range(2, ws_namen.max_row + 1):
                vn = _clean(ws_namen.cell(row=row, column=1).value)
                nn = _clean(ws_namen.cell(row=row, column=2).value)
                alias = _clean(ws_namen.cell(row=row, column=3).value)
                if alias and (vn or nn):
                    alias_by_name[f"{vn} {nn}".strip()] = alias

        for idx, sheetname in enumerate(schueler_bloetter):
            ws = wb[sheetname]
            student = self._lies_schueler(ws, sheetname)
            if student is None:
                warnungen.append(f"Blatt '{sheetname}' sieht nicht wie ein Schülerblatt aus -- übersprungen.")
                continue
            # Besuchshistorie anhand der Zeilenposition beim Einlesen
            # übernehmen (Blattreihenfolge == Ausgangszustand in "Namen");
            # ab sofort ist die Zuordnung namensgebunden (siehe Student),
            # eine geänderte Reihenfolge kann diese Werte also nicht mehr
            # der falschen Person zuschlagen.
            student.fb_besuche = besuche_by_row.get(idx + 2, [])
            student.alias = alias_by_name.get(student.voller_name, "")
            # Ein manuell eingetragenes "Letzter Besuch"-Datum, das nicht im
            # Raster auftaucht, nicht verwerfen.
            if student.letzter_besuch_fb and student.letzter_besuch_fb not in student.fb_besuche:
                student.fb_besuche.append(student.letzter_besuch_fb)
                student.fb_besuche.sort()
            if student.fb_besuche:
                student.letzter_besuch_fb = student.fb_besuche[-1]
            self.students.append(student)

        return warnungen

    def _lies_baustein_namen(self, ws: Worksheet) -> List[str]:
        namen = []
        for row in range(4, ws.max_row + 1):
            name = _clean(ws.cell(row=row, column=1).value)
            if name:
                namen.append(name)
        return namen

    def _lies_schueler(self, ws: Worksheet, sheetname: str) -> Optional[Student]:
        # Name: zuverlässig aus dem Blattnamen ableiten (Vorname zuerst,
        # so wie es die Personenblätter selbst auch tun) -- nicht aus der
        # "Namen"-Übersicht, die in anonymisierten/älteren Dateien nicht
        # zuverlässig mit den Blättern übereinstimmt.
        teile = sheetname.split(" ", 1)
        if len(teile) < 2:
            return None
        vorname, nachname = teile[0], teile[1]

        kursung = _clean(ws.cell(row=1, column=4).value)
        jahrgang_raw = ws.cell(row=1, column=6).value
        try:
            jahrgang = int(jahrgang_raw) if jahrgang_raw not in (None, "") else None
        except (ValueError, TypeError):
            jahrgang = None
        letzter_besuch = _to_date(ws.cell(row=2, column=2).value)
        hj_note = _clean(ws.cell(row=2, column=4).value)
        sonstige_deadline = _to_date(ws.cell(row=2, column=6).value)
        sonstige_bem = _clean(ws.cell(row=2, column=8).value)

        bausteine = []
        for row in range(4, ws.max_row + 1):
            werte = [ws.cell(row=row, column=c).value for c in range(1, 12)]
            name = _clean(werte[0])
            if not name:
                # Leere "Ausstehend"-Platzzeilen aus der Vorlage (Name leer,
                # aber Status vorbelegt) sind keine echten Bausteine.
                continue
            b = Baustein(
                name=name,
                status=_clean(werte[1]) or "Ausstehend",
                bausteinarbeit=_clean(werte[2]),
                lzk_datum_1=_to_date(werte[3]),
                lzk_note_1=_clean(werte[4]),
                lzk_datum_2=_to_date(werte[5]),
                lzk_note_2=_clean(werte[6]),
                halbjahr=_clean(werte[7]),
                bemerkung=_clean(werte[8]),
                lzk_bem_1=_clean(werte[9]),
                lzk_bem_2=_clean(werte[10]),
            )
            bausteine.append(b)

        return Student(
            vorname=vorname, nachname=nachname, kursung=kursung,
            jahrgangsstufe=jahrgang, letzter_besuch_fb=letzter_besuch,
            hj_note=hj_note, sonstige_deadline=sonstige_deadline,
            sonstige_deadline_bemerkung=sonstige_bem,
            bausteine=bausteine, _quellblatt=sheetname,
        )

    # ------------------------------------------------------------------
    # Neu anlegen (falls keine Datei existiert)
    # ------------------------------------------------------------------
    def neu(self, standard_bausteine: List[str]):
        self._wb = None
        self.dateipfad = None
        self.students = []
        self.vorlage_bausteine = list(standard_bausteine)

    # ------------------------------------------------------------------
    # Fristen-Logik (portiert aus dem ursprünglichen VBA-Makro
    # "NaechstesDatumZuNamenblatt")
    # ------------------------------------------------------------------
    @staticmethod
    def berechne_status(student: Student, heute: Optional[date] = None):
        """Liefert (aktueller_baustein, naechste_deadline_oder_hinweis)."""
        heute = heute or date.today()
        beste_diff = None
        ziel_datum = None
        ziel_baustein = None
        frist_noetig = False
        frist_baustein = None

        for b in student.bausteine:
            if b.status != "In Bearbeitung":
                continue
            frist_fehlt = (
                (not b.lzk_datum_1 and not b.lzk_datum_2)
                or (b.lzk_datum_1 and b.lzk_note_1 and not b.lzk_datum_2)
            )
            if frist_fehlt and frist_baustein is None:
                frist_noetig = True
                frist_baustein = b.name

            for datum, note in ((b.lzk_datum_1, b.lzk_note_1), (b.lzk_datum_2, b.lzk_note_2)):
                if datum and not note:
                    diff = abs((datum - heute).days)
                    if beste_diff is None or diff < beste_diff:
                        beste_diff = diff
                        ziel_datum = datum
                        ziel_baustein = b.name

        if ziel_datum:
            return ziel_baustein, ziel_datum
        if frist_noetig:
            return frist_baustein, "Frist vereinbaren!"
        return "", ""

    def deadline_info(self, student: Student, heute: Optional[date] = None):
        """Naechste Frist samt Anlass: (termin, anlass).

        Beruecksichtigt die LZK-Termine aus den Bausteinen UND die frei gesetzte
        Deadline der Person. Es gewinnt der fruehere Termin; ein reiner Hinweis
        ("Frist vereinbaren!") tritt hinter jeden echten Termin zurueck.

        `termin` ist ein date oder ein Hinweistext, `anlass` das, was daneben
        stehen soll -- "LZK" oder die eigene Bemerkung.
        """
        heute = heute or date.today()
        baustein, lzk_termin = self.berechne_status(student, heute)

        frei_termin = student.sonstige_deadline
        frei_anlass = (student.sonstige_deadline_bemerkung or "").strip() or "Termin"

        kandidaten = []
        if isinstance(lzk_termin, date):
            kandidaten.append((0, lzk_termin, lzk_termin, "LZK"))
        if frei_termin:
            kandidaten.append((0, frei_termin, frei_termin, frei_anlass))
        if not kandidaten and isinstance(lzk_termin, str) and lzk_termin:
            return lzk_termin, baustein or "LZK"
        if not kandidaten:
            return "", ""
        kandidaten.sort(key=lambda k: k[1])
        _, _, termin, anlass = kandidaten[0]
        return termin, anlass

    # ------------------------------------------------------------------
    # Speichern
    # ------------------------------------------------------------------
    def speichern_excel(self, pfad: str):
        if self._wb is None:
            wb = self._neue_mappe()
        else:
            wb = self._wb

        vorlage_name = next((n for n in wb.sheetnames if n.lower().startswith("vorlage")), None)
        if vorlage_name is None:
            vorlage_name = "Vorlage"
            self._baustein_blatt_erstellen(wb, vorlage_name, ist_vorlage=True)
        if vorlage_name != "Vorlage":
            wb[vorlage_name].title = "Vorlage"
            vorlage_name = "Vorlage"
        self._schreibe_vorlage(wb[vorlage_name])

        aktuelle_namen = {s.voller_name for s in self.students}

        # Blätter entfernen, die zu gelöschten Personen gehören
        for sheetname in list(wb.sheetnames):
            # "App-*"-Blaetter stammen aus dem Abgleich mit der Lerntheken-App
            # (osk_sync.py) und duerfen hier nicht wegoptimiert werden.
            if sheetname in ("Namen", "Vorlage") or sheetname.startswith("App-"):
                continue
            if sheetname not in aktuelle_namen:
                del wb[sheetname]

        # Personenblätter anlegen/aktualisieren
        for student in self.students:
            if student.voller_name not in wb.sheetnames:
                self._baustein_blatt_erstellen(wb, student.voller_name, ist_vorlage=False)
            ws = wb[student.voller_name]
            self._schreibe_schueler(ws, student)

        self._schreibe_namen_blatt(wb, self.students)

        wb.save(pfad)
        self._wb = wb
        self.dateipfad = pfad

    def _neue_mappe(self) -> Workbook:
        wb = Workbook()
        wb.remove(wb.active)
        wb.create_sheet("Namen")
        self._baustein_blatt_erstellen(wb, "Vorlage", ist_vorlage=True)
        return wb

    def _baustein_blatt_erstellen(self, wb: Workbook, name: str, ist_vorlage: bool):
        ws = wb.create_sheet(name)
        headers = ["Name:", None, "Aktuelle Kursung:", None, "Jahrgangsstufe:", None]
        ws["A1"] = "Name:"
        ws["C1"] = "Aktuelle Kursung:"
        ws["E1"] = "Jahrgangsstufe:"
        ws["A2"] = "Letzter Besuch im FB:"
        ws["C2"] = "HJ-Note:"
        kopf_labels = ["Baustein", "Status", "Bausteinarbeit", "LZK-Datum-1", "LZK-Note-1",
                       "LZK-Datum-2", "LZK-Note-2", "Halbjahr:", "Bemerkung",
                       "LZK-Bemerkung-1", "LZK-Bemerkung-2"]
        for i, label in enumerate(kopf_labels, start=1):
            ws.cell(row=3, column=i, value=label)
        for col, width in zip("ABCDEFGHIJK", [24, 16, 20, 13, 13, 13, 13, 11, 40, 26, 26]):
            ws.column_dimensions[col].width = width
        for c in range(1, 12):
            ws.cell(row=1, column=c).font = Font(bold=True)
            ws.cell(row=3, column=c).font = Font(bold=True)
        ws.freeze_panes = "A4"

    def _schreibe_vorlage(self, ws: Worksheet):
        # Bestehende Bausteinzeilen leeren, dann neu schreiben
        for row in range(4, ws.max_row + 2):
            for col in range(1, 12):
                # openpyxl ignoriert value=None beim cell()-Aufruf -- die Zelle
                # muss ueber .value geleert werden, sonst bleibt der alte Inhalt.
                ws.cell(row=row, column=col).value = None
        for i, name in enumerate(self.vorlage_bausteine):
            row = 4 + i
            ws.cell(row=row, column=1, value=name)
            ws.cell(row=row, column=2, value="Ausstehend")
        self._setze_baustein_formatierung(ws, len(self.vorlage_bausteine))

    def _schreibe_schueler(self, ws: Worksheet, student: Student):
        ws["B1"] = student.voller_name
        ws["D1"] = student.kursung
        ws["F1"] = student.jahrgangsstufe
        ws["B2"] = student.letzter_besuch_fb
        ws["D2"] = student.hj_note
        ws["E2"] = "Sonstige Deadline:"
        ws["F2"] = student.sonstige_deadline
        if student.sonstige_deadline:
            ws["F2"].number_format = "DD.MM.YYYY"
        ws["G2"] = "Anlass:"
        ws["H2"] = student.sonstige_deadline_bemerkung

        for row in range(4, ws.max_row + 2):
            for col in range(1, 12):
                # openpyxl ignoriert value=None beim cell()-Aufruf -- die Zelle
                # muss ueber .value geleert werden, sonst bleibt der alte Inhalt.
                ws.cell(row=row, column=col).value = None

        for i, b in enumerate(student.bausteine):
            row = 4 + i
            werte = [b.name, b.status, b.bausteinarbeit, b.lzk_datum_1, b.lzk_note_1,
                     b.lzk_datum_2, b.lzk_note_2, b.halbjahr, b.bemerkung,
                     b.lzk_bem_1, b.lzk_bem_2]
            for col, wert in enumerate(werte, start=1):
                ws.cell(row=row, column=col, value=wert)
            for col in (4, 6):
                if werte[col - 1] is not None:
                    ws.cell(row=row, column=col).number_format = "DD.MM.YYYY"

        self._setze_baustein_formatierung(ws, len(student.bausteine))

    def _setze_baustein_formatierung(self, ws: Worksheet, anzahl_zeilen: int):
        letzte_zeile = max(4 + anzahl_zeilen, 5)
        # Alte Validierungen/Regeln entfernen, damit sie sich nicht häufen
        ws.data_validations.dataValidation = []
        ws.conditional_formatting._cf_rules.clear()

        dv = DataValidation(type="list", formula1='"{}"'.format(",".join(STATUS_OPTIONEN)),
                             allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"B4:B{letzte_zeile}")

        for status, farbe in STATUS_FARBEN.items():
            ws.conditional_formatting.add(
                f"A4:K{letzte_zeile}",
                FormulaRule(formula=[f'$B4="{status}"'],
                            fill=PatternFill("solid", fgColor=farbe)),
            )

    def _schreibe_namen_blatt(self, wb: Workbook, students: List[Student]):
        ws = wb["Namen"]
        # Ganze bisherige Tabelle leeren -- sie wird gleich komplett aus
        # fb_besuche neu aufgebaut. WICHTIG: auch Zeile 1 ab der ersten
        # Datumsspalte. Frueher blieben dort alte Spaltenkoepfe stehen, wenn das
        # neue Raster schmaler war als das alte; beim naechsten Laden wurden sie
        # als zusaetzliche Besuchstage gelesen und die Historie bekam Dubletten.
        letzte_spalte = max(ws.max_column, 8)
        letzte_zeile = ws.max_row + 1
        for row in range(1, letzte_zeile + 1):
            erste_spalte = 1 if row > 1 else 9   # Zeile 1: Beschriftungen behalten
            for col in range(erste_spalte, letzte_spalte + 1):
                # openpyxl ignoriert value=None beim cell()-Aufruf -- die Zelle
                # muss ueber .value geleert werden, sonst bleibt der alte Inhalt.
                ws.cell(row=row, column=col).value = None

        ws["A1"] = "Vorname"
        ws["B1"] = "Name"
        ws["C1"] = "Lerntheken-Alias"
        ws["D1"] = "Aktueller Baustein"
        ws["E1"] = "Nächste Deadline"
        ws["F1"] = "Letzter Besuch im FB"
        ws["G1"] = "Kursung"
        ws["H1"] = "Anlass"
        for c in range(1, 9):
            ws.cell(row=1, column=c).font = Font(bold=True)

        # Spaltenköpfe: alle Tage, an denen irgendeine Person im FB war,
        # aufsteigend sortiert -- neue Besuchstage werden dadurch beim
        # Speichern automatisch zu neuen Spalten.
        alle_tage = sorted({d for s in students for d in s.fb_besuche})
        for j, tag in enumerate(alle_tage):
            zelle = ws.cell(row=1, column=9 + j, value=tag)
            zelle.number_format = "DD.MM.YYYY"

        for i, student in enumerate(students):
            row = 2 + i
            aktueller_baustein, _ = self.berechne_status(student)
            deadline, anlass = self.deadline_info(student)
            ws.cell(row=row, column=1, value=student.vorname)
            ws.cell(row=row, column=2, value=student.nachname)
            ws.cell(row=row, column=3, value=student.alias)
            ws.cell(row=row, column=4, value=aktueller_baustein)
            zelle = ws.cell(row=row, column=5, value=deadline if deadline != "" else None)
            if isinstance(deadline, date):
                zelle.number_format = "DD.MM.YYYY"
            ws.cell(row=row, column=8, value=anlass)
            letzter = student.fb_besuche[-1] if student.fb_besuche else student.letzter_besuch_fb
            ws.cell(row=row, column=6, value=letzter)
            ws.cell(row=row, column=7, value=student.kursung)
            # Besuchstage namensgebunden eintragen (also korrekt, auch nach
            # Sortieren/Hinzufügen/Entfernen von Personen).
            besuchte_tage = set(student.fb_besuche)
            for j, tag in enumerate(alle_tage):
                ws.cell(row=row, column=9 + j, value=1 if tag in besuchte_tage else 0)
        ws.freeze_panes = "I1"
        for col, width in zip("ABCDEFGH", [11, 11, 16, 22, 17, 16, 9, 20]):
            ws.column_dimensions[col].width = width
