# -*- coding: utf-8 -*-
"""
Testsuite für die Datenschicht der Arbeitsstände-App.

Läuft komplett eigenständig (baut sich seine Testdatei selbst) --
einfach ausführen mit: python3 test_arbeitsstaende_data.py
"""
import json
import os
import tempfile
from datetime import date

from arbeitsstaende_data import Arbeitsstaende, Baustein, Student

TMP = tempfile.mkdtemp(prefix="arbeitsstaende_test_")


def _pfad(name):
    return os.path.join(TMP, name)


def _testmappe_erzeugen(pfad):
    """Baut eine kleine, realistische Testmappe: 3 Standard-Bausteine,
    3 Personen, davon eine mit individueller Zusatzzeile."""
    az = Arbeitsstaende()
    az.neu(["Terme", "Gleichungen", "Geometrie"])

    a = Student(vorname="Anna", nachname="Adler", kursung="E", jahrgangsstufe=9, hj_note="2")
    a.bausteine = [
        Baustein(name="Terme", status="Abgeschlossen",
                 lzk_datum_1=date(2026, 2, 1), lzk_note_1="2", halbjahr="2526_1"),
        Baustein(name="Gleichungen", status="In Bearbeitung", lzk_datum_1=date(2026, 9, 10)),
        Baustein(name="Geometrie", status="Ausstehend"),
        Baustein(name="Zusatzübung Bruchrechnung", status="In Bearbeitung",
                 bemerkung="Individuell ergänzt"),
    ]

    b = Student(vorname="Ben", nachname="Bauer", kursung="G", jahrgangsstufe=8)
    b.bausteine = [Baustein(name=n, status="Ausstehend") for n in az.vorlage_bausteine]

    c = Student(vorname="Clara", nachname="Conrad", kursung="E", jahrgangsstufe=9,
                letzter_besuch_fb=date(2026, 3, 2))
    c.bausteine = [Baustein(name=n, status="Ausstehend") for n in az.vorlage_bausteine]

    az.students = [a, b, c]
    az.speichern(pfad)
    return az


def test_laden_testmappe():
    pfad = _pfad("basis.xlsx")
    _testmappe_erzeugen(pfad)

    az = Arbeitsstaende()
    warnungen = az.laden(pfad)
    assert warnungen == [], warnungen
    assert len(az.students) == 3
    assert az.vorlage_bausteine == ["Terme", "Gleichungen", "Geometrie"]

    anna = next(s for s in az.students if s.voller_name == "Anna Adler")
    assert anna.kursung == "E"
    assert anna.jahrgangsstufe == 9
    assert anna.hj_note == "2"
    terme = next(x for x in anna.bausteine if x.name == "Terme")
    assert terme.status == "Abgeschlossen"
    assert terme.lzk_note_1 == "2"
    assert terme.lzk_datum_1 == date(2026, 2, 1)
    zusatz = [x.name for x in anna.bausteine]
    assert "Zusatzübung Bruchrechnung" in zusatz, "individuelle Zeile fehlt"
    print("OK: test_laden_testmappe")


def test_speichern_ohne_aenderung_erhaelt_daten():
    pfad = _pfad("basis.xlsx")
    _testmappe_erzeugen(pfad)
    az = Arbeitsstaende()
    az.laden(pfad)
    vorher = {s.voller_name: len(s.bausteine) for s in az.students}

    pfad2 = _pfad("rt1.xlsx")
    az.speichern(pfad2)

    az2 = Arbeitsstaende()
    az2.laden(pfad2)
    nachher = {s.voller_name: len(s.bausteine) for s in az2.students}
    assert vorher == nachher, f"{vorher} vs {nachher}"

    anna1 = next(s for s in az.students if s.voller_name == "Anna Adler")
    anna2 = next(s for s in az2.students if s.voller_name == "Anna Adler")
    for x, y in zip(anna1.bausteine, anna2.bausteine):
        assert x.name == y.name and x.status == y.status and x.lzk_datum_1 == y.lzk_datum_1
    print("OK: test_speichern_ohne_aenderung_erhaelt_daten")


def test_schueler_hinzufuegen_und_speichern():
    pfad = _pfad("basis.xlsx")
    az = Arbeitsstaende()
    az.laden(pfad)
    anzahl_vorher = len(az.students)

    neu = Student(vorname="Max", nachname="Mustermann", kursung="G", jahrgangsstufe=8)
    neu.bausteine = [Baustein(name=n) for n in az.vorlage_bausteine]
    az.students.append(neu)
    pfad2 = _pfad("rt2.xlsx")
    az.speichern(pfad2)

    az2 = Arbeitsstaende()
    az2.laden(pfad2)
    assert len(az2.students) == anzahl_vorher + 1
    m = next(s for s in az2.students if s.voller_name == "Max Mustermann")
    assert m.kursung == "G" and m.jahrgangsstufe == 8
    assert len(m.bausteine) == len(az.vorlage_bausteine)
    print("OK: test_schueler_hinzufuegen_und_speichern")


def test_schueler_entfernen():
    pfad = _pfad("basis.xlsx")
    az = Arbeitsstaende()
    az.laden(pfad)
    anzahl_vorher = len(az.students)
    az.students = [s for s in az.students if s.voller_name != "Ben Bauer"]
    pfad2 = _pfad("rt3.xlsx")
    az.speichern(pfad2)

    az2 = Arbeitsstaende()
    az2.laden(pfad2)
    assert len(az2.students) == anzahl_vorher - 1
    assert "Ben Bauer" not in {s.voller_name for s in az2.students}
    assert "Anna Adler" in {s.voller_name for s in az2.students}
    print("OK: test_schueler_entfernen")


def test_fb_besuche_folgen_der_person_nicht_der_zeile():
    """Regressionstest: nach dem Entfernen einer Person dürfen FB-Besuchstage
    nicht an der falschen Person landen."""
    pfad = _pfad("mit_termine.xlsx")
    az = _testmappe_erzeugen(pfad)
    from openpyxl import load_workbook
    wb = load_workbook(pfad)
    ws = wb["Namen"]
    ws["H1"] = date(2026, 3, 1)
    ws["H1"].number_format = "DD.MM.YYYY"
    ws["H2"] = 1  # Anna
    ws["H3"] = 0  # Ben
    ws["H4"] = 1  # Clara
    wb.save(pfad)

    az2 = Arbeitsstaende()
    az2.laden(pfad)
    clara_vorher = list(next(s for s in az2.students if s.voller_name == "Clara Conrad").fb_besuche)
    # Clara hatte in der Testmappe zusätzlich ein manuell gesetztes
    # "Letzter Besuch"-Datum (2.3.), das beim Laden mit übernommen wird.
    assert clara_vorher == [date(2026, 3, 1), date(2026, 3, 2)], clara_vorher
    az2.students = [s for s in az2.students if s.voller_name != "Anna Adler"]
    pfad2 = _pfad("nach_entfernen.xlsx")
    az2.speichern(pfad2)

    az3 = Arbeitsstaende()
    az3.laden(pfad2)
    clara_nachher = next(s for s in az3.students if s.voller_name == "Clara Conrad")
    assert clara_nachher.fb_besuche == clara_vorher, \
        f"FB-Besuche sind der falschen Person gefolgt! {clara_nachher.fb_besuche} vs {clara_vorher}"
    ben_nachher = next(s for s in az3.students if s.voller_name == "Ben Bauer")
    assert ben_nachher.fb_besuche == [], ben_nachher.fb_besuche
    print("OK: test_fb_besuche_folgen_der_person_nicht_der_zeile")


def test_fb_besuch_eintragen_und_spalte_waechst():
    pfad = _pfad("basis.xlsx")
    az = Arbeitsstaende()
    az.laden(pfad)
    anna = next(s for s in az.students if s.voller_name == "Anna Adler")
    neuer_tag = date(2026, 9, 15)
    anna.fb_besuch_eintragen(neuer_tag)
    assert anna.letzter_besuch_fb == neuer_tag

    pfad2 = _pfad("rt_fb.xlsx")
    az.speichern(pfad2)

    az2 = Arbeitsstaende()
    az2.laden(pfad2)
    anna2 = next(s for s in az2.students if s.voller_name == "Anna Adler")
    assert neuer_tag in anna2.fb_besuche
    assert anna2.letzter_besuch_fb == neuer_tag
    # andere Personen duerfen dadurch keinen Besuch an diesem Tag bekommen
    ben2 = next(s for s in az2.students if s.voller_name == "Ben Bauer")
    assert neuer_tag not in ben2.fb_besuche
    print("OK: test_fb_besuch_eintragen_und_spalte_waechst")


def test_deadline_berechnung():
    az = Arbeitsstaende()
    heute = date(2026, 8, 28)

    ohne_frist = Student(vorname="A", nachname="B")
    ohne_frist.bausteine = [Baustein(name="X", status="Ausstehend")]
    assert az.berechne_status(ohne_frist, heute) == ("", "")

    mit_frist = Student(vorname="C", nachname="D")
    mit_frist.bausteine = [Baustein(name="Y", status="In Bearbeitung")]
    _, deadline = az.berechne_status(mit_frist, heute)
    assert deadline == "Frist vereinbaren!"

    mit_datum = Student(vorname="E", nachname="F")
    mit_datum.bausteine = [Baustein(name="Z", status="In Bearbeitung", lzk_datum_1=date(2026, 9, 5))]
    baustein, deadline = az.berechne_status(mit_datum, heute)
    assert deadline == date(2026, 9, 5) and baustein == "Z"
    print("OK: test_deadline_berechnung")


def test_hj_note_und_keine_phantom_bausteine():
    pfad = _pfad("basis.xlsx")
    az = Arbeitsstaende()
    az.laden(pfad)
    anna = next(s for s in az.students if s.voller_name == "Anna Adler")
    assert anna.hj_note == "2", f"HJ-Note falsch gelesen: {anna.hj_note!r}"
    for s in az.students:
        for b in s.bausteine:
            assert b.name != "", f"Leerer Baustein-Name bei {s.voller_name}"
    print("OK: test_hj_note_und_keine_phantom_bausteine")


def test_vorlage_bearbeiten_und_speichern():
    pfad = _pfad("basis.xlsx")
    az = Arbeitsstaende()
    az.laden(pfad)
    az.vorlage_bausteine.append("Neuer Standard-Baustein")
    pfad2 = _pfad("rt7.xlsx")
    az.speichern(pfad2)

    az2 = Arbeitsstaende()
    az2.laden(pfad2)
    assert "Neuer Standard-Baustein" in az2.vorlage_bausteine

    neu = Student(vorname="Vorlage", nachname="Check")
    neu.bausteine = [Baustein(name=b) for b in az2.vorlage_bausteine]
    az2.students.append(neu)
    az2.speichern(pfad2)

    az3 = Arbeitsstaende()
    az3.laden(pfad2)
    vc = next(s for s in az3.students if s.voller_name == "Vorlage Check")
    assert any(x.name == "Neuer Standard-Baustein" for x in vc.bausteine)
    print("OK: test_vorlage_bearbeiten_und_speichern")


def test_neue_leere_mappe():
    az = Arbeitsstaende()
    az.neu(["Baustein 1", "Baustein 2"])
    neu = Student(vorname="Test", nachname="Person")
    neu.bausteine = [Baustein(name=b) for b in az.vorlage_bausteine]
    az.students.append(neu)
    pfad = _pfad("neu.xlsx")
    az.speichern(pfad)

    az2 = Arbeitsstaende()
    az2.laden(pfad)
    assert len(az2.students) == 1
    assert az2.vorlage_bausteine == ["Baustein 1", "Baustein 2"]
    print("OK: test_neue_leere_mappe")


def test_json_ist_verlustfrei():
    """Excel -> JSON -> Excel muss inhaltlich dasselbe ergeben."""
    az = Arbeitsstaende()
    pfad = _pfad("basis.xlsx")
    _testmappe_erzeugen(pfad)
    az.laden(pfad)
    person = az.students[0]
    person.alias = "an.be"
    person.sonstige_deadline = date(2026, 9, 20)
    person.sonstige_deadline_bemerkung = "Referat drucken"
    person.bausteine.append(Baustein(name="Kreise", status="In Bearbeitung",
                                     halbjahr="2627_1",
                                     lzk_datum_1=date(2026, 11, 5), lzk_bem_1="nur Teil 1",
                                     lzk_datum_2=date(2027, 1, 12), lzk_note_2="3"))

    def stand(a):
        return [(s.vorname, s.nachname, s.alias, s.kursung, s.jahrgangsstufe, s.hj_note,
                 s.sonstige_deadline, s.sonstige_deadline_bemerkung,
                 tuple(s.fb_besuche), s.letzter_besuch_fb,
                 tuple((b.name, b.status, b.bausteinarbeit, b.halbjahr, b.bemerkung,
                        b.lzk_datum_1, b.lzk_note_1, b.lzk_bem_1,
                        b.lzk_datum_2, b.lzk_note_2, b.lzk_bem_2) for b in s.bausteine))
                for s in sorted(a.students, key=lambda s: s.voller_name)]

    vorher = stand(az)
    js = _pfad("arbeitsstaende.json")
    az.speichern(js)

    ausJson = Arbeitsstaende()
    assert ausJson.laden(js) == []
    assert stand(ausJson) == vorher
    assert ausJson.vorlage_bausteine == az.vorlage_bausteine

    xlsx = _pfad("aus_json.xlsx")
    ausJson.speichern_excel(xlsx)
    ausExcel = Arbeitsstaende()
    ausExcel.laden(xlsx)
    assert stand(ausExcel) == vorher
    print("OK: test_json_ist_verlustfrei")


def test_json_ist_lesbar_und_stabil():
    """Die Datei soll von Hand lesbar bleiben: Umlaute im Klartext, leere
    Felder gar nicht erst geschrieben, gleicher Stand = gleiche Datei."""
    az = Arbeitsstaende()
    az.neu(["Baustein 1"])
    s = Student(vorname="J\u00fcrgen", nachname="\u00d6ztürk", jahrgangsstufe=9)
    s.bausteine.append(Baustein(name="Kreise", status="In Bearbeitung",
                                lzk_datum_1=date(2026, 11, 5), lzk_bem_1="nur Teil 1"))
    az.students.append(s)

    js = _pfad("lesbar.json")
    az.speichern(js)
    roh = open(js, encoding="utf-8").read()
    assert "J\u00fcrgen" in roh and "\\u00fc" not in roh, "Umlaute muessen lesbar bleiben"

    daten = json.loads(roh)
    baustein = daten["personen"][0]["bausteine"][0]
    assert baustein["lzk_1"] == {"datum": "2026-11-05", "bemerkung": "nur Teil 1"}
    assert "lzk_2" not in baustein, "leere LZK gar nicht schreiben"
    assert "hj_note" not in daten["personen"][0]

    js2 = _pfad("lesbar2.json")
    Arbeitsstaende_neu = Arbeitsstaende()
    Arbeitsstaende_neu.laden(js)
    Arbeitsstaende_neu.speichern(js2)
    a, b = json.loads(roh), json.loads(open(js2, encoding="utf-8").read())
    a.pop("gespeichert_am"), b.pop("gespeichert_am")
    assert a == b
    print("OK: test_json_ist_lesbar_und_stabil")


def test_json_fremde_datei_warnt_statt_abzustuerzen():
    js = _pfad("fremd.json")
    with open(js, "w", encoding="utf-8") as f:
        json.dump({"format": "irgendwas", "version": 99,
                   "personen": [{"vorname": "Ohne", "nachname": "Bausteine"},
                                {"nachname": ""}]}, f)
    az = Arbeitsstaende()
    warnungen = az.laden(js)
    assert len(az.students) == 1
    assert len(warnungen) == 3, warnungen
    print("OK: test_json_fremde_datei_warnt_statt_abzustuerzen")


if __name__ == "__main__":
    test_laden_testmappe()
    test_speichern_ohne_aenderung_erhaelt_daten()
    test_schueler_hinzufuegen_und_speichern()
    test_schueler_entfernen()
    test_fb_besuche_folgen_der_person_nicht_der_zeile()
    test_fb_besuch_eintragen_und_spalte_waechst()
    test_deadline_berechnung()
    test_hj_note_und_keine_phantom_bausteine()
    test_vorlage_bearbeiten_und_speichern()
    test_neue_leere_mappe()
    test_json_ist_verlustfrei()
    test_json_ist_lesbar_und_stabil()
    test_json_fremde_datei_warnt_statt_abzustuerzen()
    print("\nAlle Tests erfolgreich.")
    print(f"(Testdateien lagen in {TMP})")
